import copy
import json
from pathlib import Path
from docx import Document
from docx.shared import Cm
from conftest import dossier
from litman_core.generation.bid_preparation_guide import BidPreparationGuideGenerator
from litman_core.generation.bid_preparation_guide.renderer import write_bid_preparation_guide

def _text(path: Path) -> str:
    doc = Document(path)
    return "\n".join([p.text for p in doc.paragraphs] + [cell.text for t in doc.tables for row in t.rows for cell in row.cells])

def test_bid_guide_has_five_sections_and_operational_content(tmp_path: Path) -> None:
    output = tmp_path / "guide.docx"; write_bid_preparation_guide(dossier(), output); doc, text = Document(output), _text(output)
    assert [p.text for p in doc.paragraphs if p.style.name == "Heading 2"] == ["1. Contexto", "2. Documentación que presentar", "3. Requisitos para preparar la oferta", "4. Criterios de adjudicación", "5. Pendientes y comprobación final"]
    for value in ("Memoria", "B", "Sí", "Metodología", "Explicar la metodología", "No incluir precio", "Exclusión", "Validar metodología", "Confirmar precio", "Revisar guía", "Confirmar plazo", "No consta"): assert value in text
    assert "PPT" not in text and "Anexo de fuentes y evidencias" not in text and "Calidad y limitaciones" not in text
    assert all(value not in text for value in ("Mejorar la calidad del dato", "Diseño e implantación", "Analizar las fuentes", "Plan de trabajo", "Power BI", "Jefe de proyecto"))

def test_bid_guide_uses_suggested_structure_when_imposed_sections_are_absent(tmp_path: Path) -> None:
    data = dossier(memoria_tecnica={"secciones": [], "formato": {}}, estructura_documento_tecnico=["Propuesta sugerida"]); output = tmp_path / "guide.docx"; write_bid_preparation_guide(data, output); text = _text(output)
    assert "Estructura sugerida de la memoria" in text and "Propuesta sugerida" in text

def test_bid_guide_groups_unknown_format_fields(tmp_path: Path) -> None:
    data = dossier(memoria_tecnica={"formato": {}, "portada_computa_limite": None, "indice_computa_limite": False, "anexos_computan_limite": True}); output = tmp_path / "format.docx"; write_bid_preparation_guide(data, output); text = _text(output)
    assert text.count("No consta: fuente") == 1 and "tamaño" in text and "margen superior" in text and "□ REVISAR" not in text and "□ El índice no computa en el límite." in text

def test_bid_guide_preserves_input_and_does_not_render_evidence(tmp_path: Path) -> None:
    data = dossier(); data["memoria_tecnica"]["reglas"].append({"tipo": "otra", "descripcion": "Regla especial", "consecuencia": "Revisar manualmente"}); original = copy.deepcopy(data); output = tmp_path / "complete.docx"; write_bid_preparation_guide(data, output); text = _text(output)
    assert data == original and "Regla especial" in text and "Revisar manualmente" in text and "PPT" not in text and "Página" not in text

def test_bid_guide_has_accessible_document_structure(tmp_path: Path) -> None:
    output = tmp_path / "structure.docx"; write_bid_preparation_guide(dossier(), output); doc = Document(output); section = doc.sections[0]; tolerance = Cm(.01)
    assert abs(section.page_width - Cm(21)) < tolerance and abs(section.page_height - Cm(29.7)) < tolerance and abs(section.left_margin - Cm(2)) < tolerance and abs(section.right_margin - Cm(2)) < tolerance
    assert doc.styles["Normal"].font.name == "Arial" and sum(p.style.name == "Heading 1" for p in doc.paragraphs) == 1 and not any(p._p.xpath('.//w:br[@w:type="page"]') for p in doc.paragraphs)
    assert section.header.paragraphs[0].text and section.footer._element.xpath('.//w:fldSimple[contains(@w:instr,"PAGE")]') and all(len(t.columns) <= 4 for t in doc.tables)
    for table in doc.tables: assert table.rows[0]._tr.xpath('./w:trPr/w:tblHeader') and all(row._tr.xpath('./w:trPr/w:cantSplit') for row in table.rows)
