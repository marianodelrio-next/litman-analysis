import json
import zipfile
from xml.etree import ElementTree
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from jsonschema import Draft202012Validator
from litman_core.generation import ARTIFACT_GENERATORS, resolve_artifact_generators
from litman_core.generation.service import generate_documents
from litman_core.resource_paths import RESOURCE_ROOT

from conftest import dossier


def test_generation_service_writes_traceable_artifacts(tmp_path: Path) -> None:
    dossier_path = tmp_path / "case" / "detailed-analysis" / "run-1" / "output" / "dossier.json"
    dossier_path.parent.mkdir(parents=True)
    dossier_path.write_text(json.dumps(dossier(), ensure_ascii=False), encoding="utf-8")
    run_dir = tmp_path / "generation-run"
    manifest = generate_documents(
        dossier_path, output_dir=run_dir, artifacts=["technical-document"], provider=None,
    )
    assert Path(manifest["documents"][0]).exists()
    assert manifest["generation"]["mode"] == "skeleton"
    assert manifest["content_profile"] == "skeleton-v1"
    coverage = json.loads((run_dir / "output" / "coverage.json").read_text(encoding="utf-8"))
    assert not coverage["complete"]
    assert not coverage["sections"][0]["response_complete"]
    generated_content = json.loads((run_dir / "output" / "content.json").read_text(encoding="utf-8"))
    generated_path = Path(manifest["documents"][0])
    generated_document = Document(generated_path)
    assert any(block["type"] == "subheading" for block in generated_content["sections"][0]["blocks"])
    assert any(paragraph.text != "Ejemplo de subapartado" and paragraph.style.name == "Heading 2"
               for paragraph in generated_document.paragraphs)
    with zipfile.ZipFile(generated_path) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    word = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    body_subheading = next(
        paragraph for paragraph in root.iter(word + "p")
        if paragraph.find(f"{word}pPr/{word}pStyle") is not None
        and paragraph.find(f"{word}pPr/{word}pStyle").get(word + "val") == "Heading2"
        and "Ejemplo de subapartado" not in "".join(node.text or "" for node in paragraph.iter(word + "t"))
    )
    assert body_subheading.find(f"{word}pPr/{word}numPr/{word}ilvl").get(word + "val") == "1"
    assert body_subheading.find(f"{word}pPr/{word}numPr/{word}numId").get(word + "val") == "2"
    indentation = body_subheading.find(f"{word}pPr/{word}ind")
    assert indentation.get(word + "left") == "1440"
    assert indentation.get(word + "hanging") == "360"
    toc_subheading = next(
        paragraph for paragraph in root.iter(word + "p")
        if paragraph.find(f"{word}pPr/{word}pStyle") is None
        and "1.1." in "".join(node.text or "" for node in paragraph.iter(word + "t"))
    )
    assert any(node.get(word + "val") == "22" for node in toc_subheading.iter(word + "sz"))
    assert any(node.get(word + "val") == "0" for node in toc_subheading.iter(word + "b"))
    assert any(node.get(word + "val") == "0" for node in toc_subheading.iter(word + "i"))
    assert any(node.get(word + "val") == "000000" for node in toc_subheading.iter(word + "color"))
    assert any(node.get(word + "ascii") == "Roboto Light" for node in toc_subheading.iter(word + "rFonts"))
    toc_heading = next(
        paragraph for paragraph in root.iter(word + "p")
        if paragraph.find(f"{word}pPr/{word}pStyle") is None
        and "1. Metodología" in "".join(node.text or "" for node in paragraph.iter(word + "t"))
    )
    assert any(node.get(word + "val") == "22" for node in toc_heading.iter(word + "sz"))
    assert any(node.get(word + "val") == "0" for node in toc_heading.iter(word + "b"))
    assert any(node.get(word + "val") == "000000" for node in toc_heading.iter(word + "color"))
    assert any(node.get(word + "ascii") == "Roboto Light" for node in toc_heading.iter(word + "rFonts"))
    body_heading = next(
        paragraph for paragraph in root.iter(word + "p")
        if paragraph.find(f"{word}pPr/{word}pStyle") is not None
        and paragraph.find(f"{word}pPr/{word}pStyle").get(word + "val") == "Heading1"
        and paragraph.text is None
        and "Metodología" in "".join(node.text or "" for node in paragraph.iter(word + "t"))
    )
    assert body_heading.find(f"{word}pPr/{word}numPr/{word}ilvl").get(word + "val") == "0"
    # La fuente Arial del pliego de prueba prevalece sobre Roboto de la plantilla.
    assert any(node.get(word + "ascii") == "Arial" for node in body_heading.iter(word + "rFonts"))
    assert any(node.get(word + "val") == "32" for node in body_heading.iter(word + "sz"))
    assert any(node.get(word + "val") == "4FA6DC" for node in body_heading.iter(word + "color"))
    for relative in (
        "input/document_plan.json", "input/template.yaml", "output/content.json",
        "output/pending.json", "output/coverage.json", "output/validation.json",
        "output/generation_manifest.json", "lineage.json",
    ): 
        assert (run_dir / relative).exists()
def test_artifact_registry_resolves_stable_generators() -> None:
    generators = resolve_artifact_generators(["technical-summary", "bid-preparation-guide", "technical-document", "economic-simulator"])
    assert [generator.id for generator in generators] == ["technical-summary", "bid-preparation-guide", "technical-document", "economic-simulator"]
    assert all(generator.version for generator in generators)
    assert set(ARTIFACT_GENERATORS) == {"technical-summary", "bid-preparation-guide", "technical-document", "economic-simulator"}
    assert [generator.id for generator in resolve_artifact_generators(None)] == ["technical-summary", "bid-preparation-guide", "technical-document", "economic-simulator"]


def test_unknown_artifact_is_rejected_before_generation(tmp_path: Path) -> None:
    dossier_path = tmp_path / "case" / "detailed-analysis" / "run-1" / "output" / "dossier.json"
    dossier_path.parent.mkdir(parents=True)
    dossier_path.write_text(json.dumps(dossier()), encoding="utf-8")
    try:
        generate_documents(dossier_path, output_dir=tmp_path / "run", artifacts=["not-a-document"])
    except ValueError as error:
        assert "desconocido" in str(error)
        assert "not-a-document" in str(error)
    else:
        raise AssertionError("Debe rechazar un artefacto no registrado")


def test_multiple_artifacts_are_consolidated_with_versions(tmp_path: Path) -> None:
    dossier_path = tmp_path / "case" / "detailed-analysis" / "run-1" / "output" / "dossier.json"
    dossier_path.parent.mkdir(parents=True)
    dossier_path.write_text(json.dumps(dossier()), encoding="utf-8")
    manifest = generate_documents(dossier_path, output_dir=tmp_path / "run",
                                  artifacts=["technical-summary", "bid-preparation-guide", "economic-simulator"])
    assert len(manifest["documents"]) == 3
    assert {item["id"] for item in manifest["artifacts"]} == {"technical-summary", "bid-preparation-guide", "economic-simulator"}
    assert all(item["generator_version"] for item in manifest["artifacts"])
    assert manifest["artifacts"][2]["metadata"]["template_version"] == "1.1.0"


def test_english_source_contract_fixture_generates_four_spanish_artifacts(tmp_path: Path) -> None:
    contract = json.loads((Path(__file__).parents[2] / "fixtures/spanish-output-language/english-source-contract.json").read_text())
    value = dossier()
    value["memoria_tecnica"]["idioma"] = contract["source_language"]
    value["contexto"] = contract["semantic_context"]
    value["memoria_tecnica"]["secciones"][0]["titulo"] = contract["semantic_section"]
    value["memoria_tecnica"]["secciones"][0]["subapartados"] = [{"titulo": "Enfoque técnico", "descripcion": "Descripción semántica en español", "nivel": 2, "requisitos_relacionados": [], "evidencias": []}]
    value["evidencias"] = [{"campo": "metodologia", "documento_id": "ENG-1", "pagina": 3, "texto_citado": contract["source_quote"]}]
    value["simulador_economico"]["formula_precio"]["expresion_literal"] = contract["price_formula"]
    value["simulador_economico"]["baja_anormal"]["formula_literal"] = contract["abnormal_formula"]
    assert not list(Draft202012Validator(json.loads((RESOURCE_ROOT / "schemas/detailed-analysis.schema.json").read_text())).iter_errors(value))
    dossier_path = tmp_path / "dossier.json"
    dossier_path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
    manifest = generate_documents(dossier_path, output_dir=tmp_path / "run", provider=None)
    assert {item["id"] for item in manifest["artifacts"]} == {"technical-summary", "bid-preparation-guide", "technical-document", "economic-simulator"}
    docx_text = "\n".join("\n".join(p.text for p in Document(path).paragraphs) for path in manifest["documents"] if path.endswith(".docx"))
    assert contract["semantic_context"] in docx_text
    assert "Propuesta técnica" in docx_text
    workbook = load_workbook(next(Path(path) for path in manifest["documents"] if path.endswith(".xlsx")), data_only=False)
    workbook_text = "\n".join(str(cell.value) for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row if cell.value is not None)
    assert contract["source_quote"] not in docx_text
    assert contract["price_formula"] in workbook_text
    economic_plan = json.loads((tmp_path / "run/input/economic_plan.json").read_text(encoding="utf-8"))
    assert economic_plan["abnormally_low"]["formula_literal"] == contract["abnormal_formula"]


def test_selective_summary_generation_declares_existing_output_documents(tmp_path: Path) -> None:
    dossier_path = tmp_path / "dossier.json"
    dossier_path.write_text(json.dumps(dossier()), encoding="utf-8")
    manifest = generate_documents(dossier_path, output_dir=tmp_path / "run", artifacts=["technical-summary", "bid-preparation-guide"])
    assert [item["id"] for item in manifest["artifacts"]] == ["technical-summary", "bid-preparation-guide"]
    assert {Path(path).name for path in manifest["documents"]} == {"resumen_tecnico.docx", "guia_preparacion_oferta.docx"}
    assert all(Path(path).is_file() for path in manifest["documents"])


def test_artifact_failure_identifies_generator(tmp_path: Path, monkeypatch) -> None:
    dossier_path = tmp_path / "case" / "detailed-analysis" / "run-1" / "output" / "dossier.json"
    dossier_path.parent.mkdir(parents=True)
    dossier_path.write_text(json.dumps(dossier()), encoding="utf-8")

    class BrokenSummary:
        id = "bid-preparation-guide"
        version = "test-v1"
        def generate(self, context):
            raise RuntimeError("fallo controlado")

    monkeypatch.setitem(ARTIFACT_GENERATORS, "bid-preparation-guide", BrokenSummary())
    try:
        generate_documents(dossier_path, output_dir=tmp_path / "run", artifacts=["bid-preparation-guide"])
    except ValueError as error:
        assert "bid-preparation-guide" in str(error)
        assert "fallo controlado" in str(error)
    else:
        raise AssertionError("Debe identificar el generador que falló")


def test_draft_profile_keeps_previous_deterministic_generation(tmp_path: Path) -> None:
    dossier_path = tmp_path / "case" / "detailed-analysis" / "run-1" / "output" / "dossier.json"
    dossier_path.parent.mkdir(parents=True)
    dossier_path.write_text(json.dumps(dossier(), ensure_ascii=False), encoding="utf-8")
    manifest = generate_documents(
        dossier_path, output_dir=tmp_path / "generation-run",
        artifacts=["technical-document"], provider=None, content_profile="draft-v1",
    )
    assert manifest["generation"]["mode"] == "deterministic-fallback"


def test_generated_manifest_feeds_detailed_delivery_catalog(tmp_path: Path) -> None:
    from litman_core.infrastructure.storage.common import detailed_analysis_delivery_catalog

    dossier_path = tmp_path / "dossier.json"
    dossier_path.write_text(json.dumps(dossier()), encoding="utf-8")
    run = tmp_path / "run"
    generate_documents(dossier_path, output_dir=run, provider=None)
    technical = [item for item in detailed_analysis_delivery_catalog(run) if item.role == "technical_document"]
    assert len(technical) == 1
    assert technical[0].relative_path == "propuesta_tecnica.docx"
