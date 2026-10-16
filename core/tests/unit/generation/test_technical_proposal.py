import json
import zipfile
from xml.etree import ElementTree
from pathlib import Path

import pytest

from docx import Document
from jsonschema import Draft202012Validator

from litman_core.generation.technical_proposal.content import fallback_content, skeleton_content
from litman_core.generation.technical_proposal.planner import build_document_plan
from litman_core.generation.technical_proposal.renderer import render_technical_document
from litman_core.generation.technical_proposal.generator import _validate_content
from litman_core.generation.service import generate_documents
from litman_core.cli import main
from litman_core.resource_paths import RESOURCE_ROOT

from conftest import dossier
RESOURCES = RESOURCE_ROOT


def test_detailed_analysis_fixture_matches_schema() -> None:
    schema = json.loads((RESOURCES / "schemas" / "detailed-analysis.schema.json").read_text())
    assert list(Draft202012Validator(schema).iter_errors(dossier())) == []
    assert schema["properties"]["schema_version"] == {
        "type": "string", "const": "detailed-analysis-v1",
    }


def test_detailed_analysis_schema_is_compatible_with_codex_structured_output() -> None:
    schema = json.loads((RESOURCES / "schemas" / "detailed-analysis.schema.json").read_text())

    def objects(value):
        if isinstance(value, dict):
            if "properties" in value:
                yield value
            for child in value.values():
                yield from objects(child)
        elif isinstance(value, list):
            for child in value:
                yield from objects(child)

    for object_schema in objects(schema):
        assert set(object_schema["properties"]) == set(object_schema.get("required", []))


def test_plan_preserves_tender_order_and_constraints() -> None:
    plan = build_document_plan(dossier())
    assert plan["structure_source"] == "tender"
    assert [item["title"] for item in plan["sections"]] == ["Metodología"]
    assert plan["sections"][0]["requirements"] == ["Explicar la metodología"]
    assert plan["formal_constraints"]["fuente"] == "Arial"
    assert plan["word_budget"]["total_target_words"] <= 3500
    assert plan["sections"][0]["evidence_fields"] == []


def test_generic_criterion_destination_does_not_contaminate_other_sections() -> None:
    value = dossier()
    value["criterios_adjudicacion"].append({
        "nombre": "Equipo", "puntuacion_maxima": 5, "tipo": "juicio",
        "donde_responder": "Memoria técnica", "descripcion": None,
    })
    plan = build_document_plan(value)
    assert plan["sections"][0]["criteria"] == ["Metodología"]



def test_old_dossier_gets_safe_fallback_plan() -> None:
    old = dossier()
    old.pop("memoria_tecnica")
    plan = build_document_plan(old)
    assert plan["structure_source"] == "fallback"
    assert plan["pending_decisions"]


def test_plan_uses_distinct_contextual_subsections_and_fallback_score() -> None:
    value = dossier()
    value["memoria_tecnica"]["secciones"] = []
    value["estructura_documento_tecnico"] = ["Metodología", "Equipo requerido"]
    value["criterios_adjudicacion"] = [
        {"nombre": "Metodología", "puntuacion_maxima": 5, "tipo": "juicio", "donde_responder": "Metodología", "descripcion": None},
        {"nombre": "Equipo requerido", "puntuacion_maxima": 1, "tipo": "juicio", "donde_responder": "Equipo requerido", "descripcion": None},
    ]
    plan = build_document_plan(value)
    assert [x["title"] for x in plan["sections"][0]["subsections"]] == ["Principios metodológicos", "Organización de las actividades", "Validación y control de calidad", "Seguimiento y mejora continua"]
    assert [x["title"] for x in plan["sections"][1]["subsections"]] == ["Organización del equipo", "Perfiles y responsabilidades", "Modelo de coordinación", "Disponibilidad y continuidad pendiente de confirmar"]
    assert plan["sections"][0]["display_title"] == "Metodología (5 puntos)"
    assert plan["sections"][1]["display_title"] == "Equipo requerido (1 punto)"
    assert plan["sections"][0]["score_source"] == "criterion"


def test_short_filename_preserves_acronyms_and_meaningful_terms(tmp_path: Path) -> None:
    value = dossier()
    value["cabecera"].update(nombre="Aplicaciones de Inteligencia Artificial (IA) aplicada a expedientes de Compras (CARMEN)", expediente="5100073845")
    dossier_path = tmp_path / "dossier.json"; dossier_path.write_text(json.dumps(value), encoding="utf-8")
    manifest = generate_documents(dossier_path, output_dir=tmp_path / "run", artifacts=["technical-document"], provider=None)
    assert Path(manifest["documents"][0]).name == "Propuesta técnica 5100073845 Inteligencia Artificial IA aplicada expedientes Compras CARMEN.docx"


def test_english_plan_localizes_generated_structure_and_nested_evidence() -> None:
    value = dossier()
    value["memoria_tecnica"]["idioma"] = "en"
    section = value["memoria_tecnica"]["secciones"][0]
    section.update(titulo="Methodology", descripcion=None, evidencias=[{
        "campo": "methodology", "documento_id": "PPT", "pagina": 12,
        "texto_citado": "The methodology must be described.",
    }])
    plan = build_document_plan(value)
    planned = plan["sections"][0]
    assert plan["language"] == "es"
    assert plan["source_language"] == "en"
    assert planned["display_title"] == "Metodología (20 puntos)"
    assert planned["subsections"][0]["title"] == "Principios metodológicos"
    assert planned["guidance"]["evidence_summary"][0]["document_id"] == "PPT"
    assert planned["guidance"]["evidence_summary"][0]["quoted_text"] == "The methodology must be described."
    assert "Explicar de forma verificable" in planned["guidance"]["purpose"]


def test_filename_does_not_repeat_compound_reference(tmp_path: Path) -> None:
    value = dossier()
    value["cabecera"].update(nombre="Servicio de análisis EXP-1", expediente="EXP-1")
    dossier_path = tmp_path / "dossier.json"; dossier_path.write_text(json.dumps(value), encoding="utf-8")
    manifest = generate_documents(dossier_path, output_dir=tmp_path / "run", artifacts=["technical-document"], provider=None)
    assert Path(manifest["documents"][0]).name == "Propuesta técnica EXP-1 análisis.docx"


def test_identity_pending_and_guidance_toggle_are_rendered_from_plan(tmp_path: Path) -> None:
    value = dossier(); value["cabecera"].update(nombre=None, expediente=None)
    dossier_path = tmp_path / "dossier.json"; dossier_path.write_text(json.dumps(value), encoding="utf-8")
    run = tmp_path / "run"
    generate_documents(dossier_path, output_dir=run, artifacts=["technical-document"], provider=None)
    pending = json.loads((run / "output/pending.json").read_text(encoding="utf-8"))
    assert sum("nombre ni expediente" in item["description"] for item in pending) == 1
    manifest = generate_documents(dossier_path, output_dir=tmp_path / "clean", artifacts=["technical-document"],
                                  provider=None, include_editorial_guidance=False)
    clean = Document(manifest["documents"][0])
    assert not any("Guía de redacción" in p.text for p in clean.paragraphs)
    assert all("EEF5F8" not in str(table._tbl.xml) for table in clean.tables)


def test_render_uses_clean_template_and_keeps_dynamic_section(tmp_path: Path) -> None:
    plan = build_document_plan(dossier())
    content = fallback_content(dossier(), plan)
    content["sections"][0]["blocks"].append({
        "type": "table", "text": None, "items": [],
        "headers": ["Fase", "Plazo"], "rows": [["Diseño", "2 semanas"]],
    })
    output = tmp_path / "proposal.docx"
    validation = render_technical_document(
        RESOURCES / "templates" / "technical-proposal" / "next-standard-v1" / "template.docx",
        content, plan, output,
    )
    assert validation["valid"]
    text = "\n".join(item.text for item in Document(output).paragraphs)
    assert "1. Metodología" in text
    assert "OCS" not in text
    assert "[[DOCUMENT_BODY]]" not in text
    assert "Propuesta técnica" in text
    assert "Índice" in text
    assert "Technical proposal" not in text
    assert "1. Metodología" in text
    assert validation["page_limit_validation"] == "not_performed"
    rendered = Document(output)
    assert [paragraph.text for paragraph in rendered.paragraphs[:8]] == [
        "", "", "", "", "", "", "Propuesta técnica", "Servicio de analítica\nExpte. EXP-1",
    ]
    assert rendered.tables
    shading = rendered.tables[0].rows[0].cells[0]._tc.tcPr.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd"
    )
    assert shading is not None
    assert shading.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill") == "EEF5F8"

    source = RESOURCES / "templates" / "technical-proposal" / "next-standard-v1" / "template.docx"
    with zipfile.ZipFile(source) as template_archive, zipfile.ZipFile(output) as output_archive:
        assert output_archive.read("word/header1.xml") == template_archive.read("word/header1.xml")
        assert output_archive.read("word/media/image25.png") == template_archive.read("word/media/image25.png")
        root = ElementTree.fromstring(output_archive.read("word/document.xml"))
    word = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    toc = [node.text or "" for node in root.iter(word + "instrText") if "TOC" in (node.text or "")]
    pagerefs = [node for node in root.iter(word + "fldSimple") if "PAGEREF" in node.get(word + "instr", "")]
    bookmarks = [node.get(word + "name") for node in root.iter(word + "bookmarkStart")
                 if (node.get(word + "name") or "").startswith("litman_core_")]
    blue_runs = [node for node in root.iter(word + "color") if node.get(word + "val") == "4FA6DC"]
    assert toc
    assert len(pagerefs) >= 1
    assert len(bookmarks) >= 1
    assert blue_runs


def test_render_keeps_spanish_interface_for_historical_english_plan(tmp_path: Path) -> None:
    plan = build_document_plan(dossier())
    plan["language"] = "en"
    content = fallback_content(dossier(), plan)
    output = tmp_path / "proposal-en.docx"
    render_technical_document(
        RESOURCES / "templates" / "technical-proposal" / "next-standard-v1" / "template.docx",
        content, plan, output,
    )
    text = "\n".join(item.text for item in Document(output).paragraphs)
    assert "Propuesta técnica" in text
    assert "Índice" in text
    assert "Technical proposal" not in text


def test_generated_spanish_headings_keep_section_and_subsection_ids(tmp_path: Path) -> None:
    value = dossier()
    value["memoria_tecnica"]["secciones"][0]["titulo"] = "Methodology"
    plan = build_document_plan(value)
    content = fallback_content(value, plan)
    content["sections"][0]["title"] = "Metodología"
    subheading = next(block for block in content["sections"][0]["blocks"] if block["type"] == "subheading")
    subheading["text"] = "Principios metodológicos"
    assert _validate_content(content, plan)["valid"]
    output = tmp_path / "proposal.docx"
    render_technical_document(RESOURCES / "templates" / "technical-proposal" / "next-standard-v1" / "template.docx", content, plan, output)
    text = "\n".join(item.text for item in Document(output).paragraphs)
    assert "1. Metodología" in text


def test_standard_generation_hides_english_tender_structure_from_proposal(tmp_path: Path) -> None:
    value = dossier()
    value["memoria_tecnica"]["idioma"] = "en"
    value["memoria_tecnica"]["secciones"] = [{
        "orden": 1, "titulo": "Methodology", "descripcion": "Describe the delivery approach.",
        "obligatoria": True, "puntuacion": 20, "limite_paginas": 2,
        "requisitos_relacionados": [], "evidencias": [],
        "subapartados": [{"titulo": "Technical approach", "descripcion": "Explain the approach.", "nivel": 2,
                            "requisitos_relacionados": [], "evidencias": []}],
    }]
    dossier_path = tmp_path / "dossier.json"
    dossier_path.write_text(json.dumps(value), encoding="utf-8")
    manifest = generate_documents(dossier_path, output_dir=tmp_path / "run", artifacts=["technical-document"], provider=None)
    plan = json.loads((tmp_path / "run/input/document_plan.json").read_text(encoding="utf-8"))
    text = "\n".join(paragraph.text for paragraph in Document(manifest["documents"][0]).paragraphs)
    assert plan["sections"][0]["source_title"] == "Methodology"
    assert plan["sections"][0]["title"] == "Metodología"
    assert plan["sections"][0]["subsections"][0]["source_title"] == "Technical approach"
    assert "Methodology" not in text
    assert "Technical approach" not in text
    assert "Describe the delivery approach." not in text
    assert "Explain the approach." not in text
    assert "Metodología" in text
    assert "Enfoque técnico" in text


def test_spanish_titles_and_descriptions_are_preserved_without_heuristics() -> None:
    value = dossier()
    section = value["memoria_tecnica"]["secciones"][0]
    section.update(titulo="Plan de implantación", descripcion="Descripción española exacta")
    section["subapartados"] = [{"titulo": "Preparación", "descripcion": "Detalle español exacto", "nivel": 2,
                                 "requisitos_relacionados": [], "evidencias": []}]
    plan = build_document_plan(value)
    assert plan["sections"][0]["title"] == "Plan de implantación"
    assert plan["sections"][0]["purpose"] == "Descripción española exacta"
    assert plan["sections"][0]["guidance"]["purpose"].startswith("Descripción española exacta")
    assert plan["sections"][0]["subsections"][0]["title"] == "Preparación"
    assert plan["sections"][0]["subsections"][0]["purpose"] == "Detalle español exacto"


def test_declared_english_tender_preserves_unknown_translated_spanish_titles() -> None:
    value = dossier()
    value["memoria_tecnica"]["idioma"] = "en"
    value["memoria_tecnica"]["secciones"] = [
        {"orden": 1, "titulo": "Estrategia de implantación", "descripcion": "Descripción española exacta", "obligatoria": True,
         "puntuacion": None, "limite_paginas": None, "requisitos_relacionados": [], "evidencias": [], "subapartados": []},
        {"orden": 2, "titulo": "Modelo operativo", "descripcion": "Otra descripción española", "obligatoria": True,
         "puntuacion": None, "limite_paginas": None, "requisitos_relacionados": [], "evidencias": [], "subapartados": []},
    ]
    plan = build_document_plan(value)
    assert [section["title"] for section in plan["sections"]] == ["Estrategia de implantación", "Modelo operativo"]
    assert plan["sections"][0]["purpose"] == "Descripción española exacta"


def test_cli_rejects_legacy_english_output_option() -> None:
    with pytest.raises(SystemExit, match="2"):
        main(["analyze", "case", "--output-language", "en"])


def test_clean_template_contains_no_historical_body() -> None:
    template = RESOURCES / "templates" / "technical-proposal" / "next-standard-v1" / "template.docx"
    with zipfile.ZipFile(template) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    for term in ("OCS", "DDnA", "Self-Serve Data", "Aena T4"):
        assert term not in xml
    for marker in ("[[COVER_LABEL]]", "[[COVER_TITLE]]", "[[INDEX_TITLE]]", "[[INDEX_ENTRIES]]"):
        assert marker in xml


def test_skeleton_content_has_contextual_subsections_and_editorial_guidance() -> None:
    plan = build_document_plan(dossier())
    content = skeleton_content(dossier(), plan)
    section = content["sections"][0]
    rendered_words = sum(len(item.split()) for block in section["blocks"] for item in block["items"])
    assert rendered_words < 100
    assert section["requirements_covered"] == []
    assert section["pending"]
    subheading = next(block for block in section["blocks"] if block["type"] == "subheading")
    assert subheading["text"] != "Ejemplo de subapartado"
    assert subheading["level"] == 2
    assert subheading["id"].startswith(section["id"])
    assert plan["sections"][0]["guidance"]["purpose"]
    assert content["warnings"] == ["Documento esqueleto: no contiene una propuesta técnica redactada."]
