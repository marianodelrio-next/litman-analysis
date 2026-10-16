import copy
import json
from pathlib import Path
from docx import Document
from docx.shared import Cm
from conftest import dossier
from litman_core.generation.technical_summary import TechnicalSummaryGenerator
from litman_core.generation.technical_summary.renderer import write_technical_summary

def _text(path: Path) -> str:
    doc = Document(path)
    return "\n".join([p.text for p in doc.paragraphs] + [cell.text for table in doc.tables for row in table.rows for cell in row.cells])

def test_summary_has_one_title_and_five_sections_in_order(tmp_path: Path) -> None:
    output = tmp_path / "summary.docx"; write_technical_summary(dossier(), output)
    doc = Document(output); text = _text(output)
    assert [p.text for p in doc.paragraphs if p.style.name == "Heading 2"] == ["1. Contexto", "2. Objetivos", "3. Fases y entregables", "4. Tecnologías y recursos", "5. Aspectos clave"]
    assert sum(p.style.name == "Heading 1" for p in doc.paragraphs) == 1
    assert "Guía para el revisor técnico" not in text and "Comprensión del proyecto" not in text and "None" not in text

def test_summary_preserves_compact_content_and_no_forced_cover_break(tmp_path: Path) -> None:
    data = dossier(contexto="Contexto " + "x" * 500, objetivos=["Objetivo " + str(i) + " x" * 100 for i in range(5)])
    output = tmp_path / "summary.docx"; write_technical_summary(data, output); doc = Document(output); text = _text(output)
    assert "Objetivo 0" in text and "Objetivo 4" in text and not any(p._p.xpath('.//w:br[@w:type="page"]') for p in doc.paragraphs)

def test_summary_has_explicit_empty_states_and_conditional_alert(tmp_path: Path) -> None:
    data = dossier(objetivos=[], fases=[], tecnologias=[], equipo=[], requisitos_oferta=[], criterios_adjudicacion=[], calidad={"limitaciones": []})
    output = tmp_path / "empty.docx"; write_technical_summary(data, output); text = _text(output)
    for value in ("No consta información sobre objetivos.", "No consta información sobre fases.", "No constan requisitos de la oferta.", "No constan criterios de adjudicación.", "No consta"): assert value in text
    assert "ALERTA: LIMITACIONES" not in text
    output = tmp_path / "alert.docx"; write_technical_summary(dossier(), output); assert "ALERTA: LIMITACIONES DEL ANÁLISIS" in _text(output)

def test_summary_has_horizontal_phase_table_and_accessible_layout(tmp_path: Path) -> None:
    output = tmp_path / "structure.docx"; data = dossier(); original = copy.deepcopy(data); write_technical_summary(data, output)
    doc = Document(output); section = doc.sections[0]; tolerance = Cm(.01)
    assert data == original and abs(section.page_width - Cm(21)) < tolerance and abs(section.page_height - Cm(29.7)) < tolerance
    assert abs(section.left_margin - Cm(2)) < tolerance and abs(section.right_margin - Cm(2)) < tolerance and doc.styles["Normal"].font.name == "Arial"
    assert section.header.paragraphs[0].text and section.footer._element.xpath('.//w:fldSimple[contains(@w:instr,"PAGE")]')
    for table in doc.tables:
        assert table.rows[0]._tr.xpath('./w:trPr/w:tblHeader') and all(row._tr.xpath('./w:trPr/w:cantSplit') for row in table.rows)

def test_summary_header_complete_and_partial_has_no_none(tmp_path: Path) -> None:
    complete = dossier()
    complete["cabecera"].update({"fecha_limite": "2026-10-20", "presupuesto_sin_iva": 1200, "moneda": "EUR"})
    partial = dossier()
    partial["cabecera"].update({"fecha_limite": None, "presupuesto_sin_iva": 1200, "moneda": None})
    for index, data in enumerate((complete, partial)):
        output = tmp_path / f"header-{index}.docx"; write_technical_summary(data, output)
        text = _text(output)
        assert "EXP-1" in text and "Servicio de analítica" in text
        assert "None" not in text
    assert "1200 EUR" in _text(tmp_path / "header-0.docx")
    assert "1200" in _text(tmp_path / "header-1.docx")

def test_summary_handles_phases_with_missing_fields(tmp_path: Path) -> None:
    data = dossier(fases=[{}, {"nombre": "Fase incompleta", "entregables": []}])
    output = tmp_path / "missing-phase-fields.docx"; write_technical_summary(data, output)
    text = _text(output)
    assert "Fase incompleta" in text and "No consta" in text and "None" not in text

def test_summary_preserves_all_phase_deliverables_and_abbreviates_only_description(tmp_path: Path) -> None:
    description = "Descripción extensa " + "x" * 500
    data = dossier(fases=[{"nombre": "Ejecución", "descripcion": "Actividad", "plazo": "1 mes", "entregables": [
        "Entregable crítico A", "Entregable crítico B", {"nombre": "Entregable crítico C", "descripcion": description},
    ]}])
    output = tmp_path / "deliverables.docx"; write_technical_summary(data, output)
    text = _text(output)
    assert all(value in text for value in ("Entregable crítico A", "Entregable crítico B", "Entregable crítico C"))
    assert len(description) > 180 and "Descripción extensa" in text and description not in text
    assert "None" not in text
