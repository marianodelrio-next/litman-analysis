"""Deterministic text extraction from published artifacts (no LLM, no network).

The returned text is document content for a person or an assistant to read; it
is never interpreted as instructions by the harness.
"""
from __future__ import annotations

from pathlib import Path

from ..errors import HarnessError

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
SUPPORTED_MIME_TYPES = (PDF_MIME, DOCX_MIME)


def extract_text(path: Path, mime_type: str) -> str:
    """PDF text page by page (``pypdf``); DOCX paragraphs and table cells (``python-docx``)."""
    if mime_type not in SUPPORTED_MIME_TYPES:
        raise HarnessError("unsupported_artifact", "Sólo se puede leer texto de artefactos PDF o DOCX",
                           details={"mime_type": mime_type})
    try:
        return _pdf_text(path) if mime_type == PDF_MIME else _docx_text(path)
    except HarnessError:
        raise
    except Exception as exc:  # noqa: BLE001 - any parser failure is an unreadable artifact
        raise HarnessError("unreadable_artifact", "No se pudo extraer el texto del artefacto",
                           details={"error_type": type(exc).__name__}) from exc


def _pdf_text(path: Path) -> str:
    from pypdf import PdfReader
    pages = [(page.extract_text() or "").strip() for page in PdfReader(str(path)).pages]
    return "\n\n".join(page for page in pages if page)


def _docx_text(path: Path) -> str:
    from docx import Document
    document = Document(str(path))
    lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)
