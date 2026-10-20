"""Navigation and general reading of the files of one dossier: list, search, inspect, read and comments.

Every query proves that the file is inside the connected tree, and not at or below ``_LitMan``, before
any content call.  Content is returned as text for the assistant to read, never as instructions.
"""
from __future__ import annotations

import datetime
from pathlib import Path
import tempfile
from typing import Any

from ..adapters.documents import DOCX_MIME, PDF_MIME, extract_text
from ..adapters.sheets import FORMATTED_VALUE, SheetsGateway
from ..errors import HarnessError
from .documents import CONTENT_NOTE, DocumentService
from .manifest import Manifest
from .membership import belongs_to_root, in_litman
from .models import (FOLDER_MIME, GOOGLE_DOC_MIME, GOOGLE_SHEET_MIME, GOOGLE_SLIDES_MIME, XLSX_MIME, Resource)
from .service import ExpedienteService
from .summary import PAGE_SIZE
from .tree import LITMAN_FOLDER, normalized_name, walk_tree

KINDS = ("folder", "google_doc", "google_sheet", "google_slides", "pdf", "docx", "xlsx", "text", "other")
MAX_FILE_BYTES = 30 * 1024 * 1024
MAX_CELLS = 200_000
SEARCH_LIMIT = 50
TEXT_MIMES = ("text/plain", "text/markdown", "text/x-markdown", "text/csv")
TEXT_SUFFIXES = (".md", ".txt", ".csv")
_BY_MIME = {FOLDER_MIME: "folder", GOOGLE_DOC_MIME: "google_doc", GOOGLE_SHEET_MIME: "google_sheet",
            GOOGLE_SLIDES_MIME: "google_slides", PDF_MIME: "pdf", DOCX_MIME: "docx", XLSX_MIME: "xlsx"}
_CAPABILITY_FLAGS = (("can_edit", "canEdit"), ("can_rename", "canRename"), ("can_move", "canMoveItemWithinDrive"),
                     ("can_trash", "canTrash"), ("can_copy", "canCopy"), ("can_download", "canDownload"),
                     ("can_comment", "canComment"))


def file_kind(item: Resource) -> str:
    """The business kind of a Drive file; plain text is recognised by MIME type or by file extension."""
    kind = _BY_MIME.get(item.mime_type)
    if kind:
        return kind
    if item.mime_type in TEXT_MIMES or item.name.lower().endswith(TEXT_SUFFIXES):
        return "text"
    return "other"


def _too_large(message: str, **details: Any) -> HarnessError:
    return HarnessError("file_too_large", message, details=details)


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        value = value.isoformat()
    return str(value).replace("\t", " ").replace("\r", " ").replace("\n", " ")


def render_sheet(title: str, rows: Any) -> str:
    """``## Hoja: <title>`` and one ``<row number>\\t<cell>…`` line for each non-empty row."""
    lines = [f"## Hoja: {title}"]
    for number, row in enumerate(rows, start=1):
        cells = [_cell(value) for value in row]
        while cells and not cells[-1]:
            cells.pop()
        if cells:
            lines.append("\t".join([str(number), *cells]))
    return "\n".join(lines)


def _quoted_range(title: str) -> str:
    return "'" + title.replace("'", "''") + "'"


class FileService:
    def __init__(self, dossier: ExpedienteService, documents: DocumentService, sheets: SheetsGateway):
        self.dossier = dossier
        self.drive = dossier.drive
        self.documents = documents
        self.sheets = sheets

    # -- navigation ------------------------------------------------------------------------------

    def list_folder(self, folder_id: str, parent_folder_id: str | None = None) -> dict[str, Any]:
        """The children of the root or of a folder of the tree; ``_LitMan`` is never listed."""
        root = self.dossier.root(folder_id)
        folder = self._folder(root, parent_folder_id or root.id)
        roles = self._roles(root)
        children = [item for item in self.drive.children(folder.id, root.drive_id) if item.name != LITMAN_FOLDER]
        items = [self._summary(item, roles) for item in sorted(children, key=lambda item: (item.name, item.id))]
        return {"folder_id": root.id, "parent_folder_id": folder.id, "items": items, "writes": False}

    def search_in_dossier(self, folder_id: str, query: str, kind: str | None = None) -> dict[str, Any]:
        """Files whose name contains the query, ignoring case and accents, below the root and outside ``_LitMan``."""
        needle = normalized_name(query) if isinstance(query, str) else ""
        if not needle:
            raise HarnessError("invalid_request", "query debe contener texto")
        if kind is not None and kind not in KINDS:
            raise HarnessError("invalid_request", "Tipo de archivo no admitido", details={"allowed": list(KINDS)})
        root = self.dossier.root(folder_id)
        roles = self._roles(root)
        tree, walk_truncated = walk_tree(self.drive, root)
        found = [entry for entry in tree if entry.resource.name != LITMAN_FOLDER
                 and needle in normalized_name(entry.resource.name)
                 and (kind is None or file_kind(entry.resource) == kind)]
        items = [{**self._summary(entry.resource, roles), "path": entry.path} for entry in found[:SEARCH_LIMIT]]
        return {"folder_id": root.id, "query": query, "items": items,
                "truncated": walk_truncated or len(found) > SEARCH_LIMIT, "writes": False}

    def get_file_info(self, folder_id: str, file_id: str) -> dict[str, Any]:
        root = self.dossier.root(folder_id)
        item = self._resolve_in_tree(root, file_id)
        info = {**self._summary(item, self._roles(root)), "path": self._path(root, item)}
        info.update({flag: bool(item.capabilities.get(capability)) for flag, capability in _CAPABILITY_FLAGS})
        return {"folder_id": root.id, **info, "writes": False}

    # -- reading ---------------------------------------------------------------------------------

    def read_file(self, folder_id: str, file_id: str, offset: int = 0) -> dict[str, Any]:
        """One page of the text of any supported file of the dossier; writes nothing."""
        if not isinstance(offset, int) or isinstance(offset, bool) or offset < 0:
            raise HarnessError("invalid_request", "offset debe ser un entero mayor o igual que 0")
        root = self.dossier.root(folder_id)
        item = self._resolve_in_tree(root, file_id)
        kind = file_kind(item)
        extra: dict[str, Any] = {}
        if kind == "google_doc":
            page = self.documents.read(root.id, item.id, offset)
            text, extra = page["text"], {"revision_id": page["revision_id"]}
            total, next_offset = page["characters"], page["next_offset"]
        else:
            full = self._text(item, kind)
            end = offset + PAGE_SIZE
            text, total, next_offset = full[offset:end], len(full), end if end < len(full) else None
        return {"folder_id": root.id, "file_id": item.id, "name": item.name, "mime_type": item.mime_type,
                "kind": kind, "text": text, "offset": offset, "next_offset": next_offset, "characters": total,
                **extra, "note": CONTENT_NOTE, "writes": False}

    def list_comments(self, folder_id: str, file_id: str, include_resolved: bool = False) -> dict[str, Any]:
        root = self.dossier.root(folder_id)
        item = self._resolve_in_tree(root, file_id)
        if item.mime_type == FOLDER_MIME:
            raise HarnessError("unsupported_document", "Una carpeta no tiene comentarios",
                               details={"mime_type": item.mime_type})
        comments = [_comment(raw) for raw in self.drive.comments(item.id)
                    if include_resolved or not raw.get("resolved")]
        return {"folder_id": root.id, "file_id": item.id, "name": item.name, "include_resolved": include_resolved,
                "comments": comments, "note": CONTENT_NOTE, "writes": False}

    # -- helpers ---------------------------------------------------------------------------------

    def _resolve_in_tree(self, root: Resource, file_id: str) -> Resource:
        """A file inside the tree and not at or below ``_LitMan``, proved before reading its content."""
        item = self.drive.get(file_id)
        if not belongs_to_root(self.drive, root, item) or in_litman(self.drive, root, item):
            raise HarnessError("document_outside_tree", "El archivo no está dentro de la licitación conectada",
                               details={"file_id": file_id})
        return item

    def _folder(self, root: Resource, folder_id: str) -> Resource:
        """The root or a folder of the tree that is neither ``_LitMan`` nor below it."""
        folder = root if folder_id == root.id else self.drive.get(folder_id)
        inside = folder.id == root.id or (belongs_to_root(self.drive, root, folder)
                                          and not in_litman(self.drive, root, folder))
        if folder.mime_type != FOLDER_MIME or not inside:
            raise HarnessError("invalid_parent", "La carpeta debe ser una carpeta de la licitación fuera "
                               f"de {LITMAN_FOLDER}", details={"parent_folder_id": folder_id})
        return folder

    def _path(self, root: Resource, item: Resource) -> str:
        names = [item.name]
        current, seen = item, {item.id}
        while root.id not in current.parents:
            parent = next((candidate for candidate in map(self.drive.get, current.parents)
                           if candidate.id not in seen and belongs_to_root(self.drive, root, candidate)), None)
            if parent is None:
                break
            names.insert(0, parent.name)
            seen.add(parent.id)
            current = parent
        return "/".join(names)

    def _roles(self, root: Resource) -> dict[str, str]:
        """File ID → role in an adopted dossier; empty when the dossier has no manifest."""
        manifest: Manifest | None = self.dossier.load_state(root).manifest
        if manifest is None:
            return {}
        roles = {identifier: f"folder:{key}" for key, identifier in manifest.folders.items() if identifier}
        roles.update({analysis["folder_id"]: "analysis_folder" for analysis in manifest.analyses
                      if analysis.get("folder_id")})
        if manifest.context_file_id:
            roles[manifest.context_file_id] = "context"
        if manifest.live_offer:
            roles[manifest.live_offer["document_id"]] = "live_offer"
        roles.update(manifest.summary_sources())
        return roles

    @staticmethod
    def _summary(item: Resource, roles: dict[str, str]) -> dict[str, Any]:
        return {"id": item.id, "name": item.name, "mime_type": item.mime_type, "kind": file_kind(item),
                "size": None if item.mime_type.startswith("application/vnd.google-apps.") else item.size,
                "modified_time": item.modified_time, "web_view_link": item.web_view_link,
                "role": roles.get(item.id)}

    def _text(self, item: Resource, kind: str) -> str:
        if kind == "google_sheet":
            return self._sheet_text(item)
        if kind not in ("pdf", "docx", "xlsx", "text", "google_slides"):
            raise HarnessError("unsupported_document", "Este tipo de archivo no se puede leer como texto",
                               details={"mime_type": item.mime_type, "kind": kind})
        if item.size is not None and item.size > MAX_FILE_BYTES:
            raise _too_large("El archivo supera los 30 MB que LitMan puede leer", size=item.size)
        with tempfile.TemporaryDirectory(prefix="litman-file-") as directory:
            path = Path(directory) / ("file.xlsx" if kind == "xlsx" else "file")  # openpyxl checks the extension
            if kind == "google_slides":
                self.drive.export(item.id, "text/plain", path)
            else:
                self.drive.download(item.id, path)
            if kind in ("pdf", "docx"):
                return extract_text(path, item.mime_type)
            if kind == "xlsx":
                return _workbook_text(path)
            return path.read_bytes().decode("utf-8", errors="replace")

    def _sheet_text(self, item: Resource) -> str:
        sheets = (self.sheets.get(item.id).get("sheets") or [])
        properties = [entry.get("properties") or {} for entry in sheets]
        cells = sum(int((grid := prop.get("gridProperties") or {}).get("rowCount") or 0)
                    * int(grid.get("columnCount") or 0) for prop in properties)
        if cells > MAX_CELLS:
            raise _too_large("La hoja de cálculo supera las 200.000 celdas que LitMan puede leer", cells=cells)
        titles = [str(prop.get("title") or "") for prop in properties]
        if not titles:
            return ""
        ranges = self.sheets.values(item.id, [_quoted_range(title) for title in titles], FORMATTED_VALUE)
        return "\n".join(render_sheet(title, (ranges[number].get("values") or []) if number < len(ranges) else [])
                         for number, title in enumerate(titles))


def _workbook_text(path: Path) -> str:
    """Cached values of every sheet of an XLSX, read without loading the whole workbook."""
    from openpyxl import load_workbook
    try:
        workbook = load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any parser failure is an unreadable file
        raise HarnessError("unreadable_document", "No se pudo leer el archivo XLSX",
                           details={"error_type": type(exc).__name__}) from exc
    try:
        counted = 0
        sections = []
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                counted += len(row)
                if counted > MAX_CELLS:
                    raise _too_large("La hoja de cálculo supera las 200.000 celdas que LitMan puede leer")
                rows.append(row)
            sections.append(render_sheet(sheet.title, rows))
        return "\n".join(sections)
    finally:
        workbook.close()


def _comment(raw: dict[str, Any]) -> dict[str, Any]:
    quoted = raw.get("quotedFileContent")
    return {"id": raw.get("id"), "content": raw.get("content"), "author": _author(raw),
            "created_time": raw.get("createdTime"), "resolved": bool(raw.get("resolved")),
            "quoted_text": quoted.get("value") if isinstance(quoted, dict) else None,
            "replies": [{"id": reply.get("id"), "content": reply.get("content"), "author": _author(reply),
                         "created_time": reply.get("createdTime"), "action": reply.get("action")}
                        for reply in raw.get("replies") or [] if not reply.get("deleted")]}


def _author(raw: dict[str, Any]) -> str | None:
    author = raw.get("author")
    return author.get("displayName") if isinstance(author, dict) else None
