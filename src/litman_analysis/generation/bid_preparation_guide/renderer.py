from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document

from ..docx_helpers import (MISSING, add_table, add_table_or_missing, amount_value,
                            boolean_value, display_value, join_values, named_description,
                            truncate_value)


def write_bid_preparation_guide(data: dict[str, Any], path: Path) -> None:
    doc = Document()
    _title_and_header(doc, "Guía de preparación de la oferta", data.get("cabecera") or {})
    _offer_documents(doc, data.get("documentos_oferta"))
    _technical_memory(doc, data)
    _criteria(doc, data.get("criterios_adjudicacion"), "Criterios de adjudicación")
    _requirements(doc, data.get("requisitos_oferta"), "Requisitos de la oferta")
    _pending(doc, data.get("pendientes"))
    _quality(doc, data.get("calidad"))
    _evidence(doc, data.get("evidencias"))
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def _title_and_header(doc: Document, title: str, header: dict[str, Any]) -> None:
    doc.add_heading(title, 0)
    add_table(doc, ["Campo", "Valor"], [
        ("Oportunidad", display_value(header.get("nombre"))), ("Expediente", display_value(header.get("expediente"))),
        ("Organismo", display_value(header.get("organismo"))), ("Fecha límite", display_value(header.get("fecha_limite"))),
        ("Presupuesto sin IVA", amount_value(header.get("presupuesto_sin_iva"), header.get("moneda"))),
        ("Moneda", display_value(header.get("moneda"))),
    ])


def _offer_documents(doc: Document, values: Any) -> None:
    doc.add_heading("Documentos que presentar con la oferta", 1)
    doc.add_paragraph("Esta relación corresponde a la presentación de la oferta; no a los entregables del contrato.")
    add_table_or_missing(doc, ["Documento", "Sobre", "Obligatorio", "Descripción"], values,
        lambda item: (display_value(item.get("nombre")), display_value(item.get("sobre")), boolean_value(item.get("obligatorio")), display_value(item.get("descripcion"))))


def _technical_memory(doc: Document, data: dict[str, Any]) -> None:
    memory = data.get("memoria_tecnica") if isinstance(data.get("memoria_tecnica"), dict) else {}
    doc.add_heading("Contrato de la memoria técnica", 1)
    add_table(doc, ["Campo", "Valor"], [("Requerida", boolean_value(memory.get("requerida"))),
        ("Nombre del documento", display_value(memory.get("nombre_documento"))), ("Idioma", display_value(memory.get("idioma"))),
        ("Estructura impuesta", boolean_value(memory.get("estructura_impuesta")))])
    doc.add_heading("Estructura de la memoria técnica", 1)
    sections = memory.get("secciones")
    if isinstance(sections, list) and sections:
        valid = [item for item in sections if isinstance(item, dict)]
        rows = [(display_value(item.get("orden")), display_value(item.get("titulo")), display_value(item.get("descripcion")), boolean_value(item.get("obligatoria")),
                 display_value(item.get("puntuacion")), display_value(item.get("limite_paginas")), join_values(item.get("requisitos_relacionados")))
                for item in sorted(valid, key=lambda item: (item.get("orden") is None, item.get("orden", 0)))]
        add_table(doc, ["Orden", "Título", "Descripción", "Obligatoria", "Puntuación", "Límite páginas", "Requisitos relacionados"], rows or [tuple(MISSING for _ in range(7))])
    else:
        suggested = data.get("estructura_documento_tecnico")
        if isinstance(suggested, list) and suggested:
            doc.add_paragraph("Estructura sugerida por el análisis; no es una estructura obligatoria del pliego.")
            for item in suggested:
                doc.add_paragraph(named_description(item), style="List Bullet")
        else:
            doc.add_paragraph("No se ha localizado una estructura de la memoria técnica.")
    doc.add_heading("Formato y límites de la memoria", 1)
    fmt = memory.get("formato") if isinstance(memory.get("formato"), dict) else {}
    add_table(doc, ["Campo", "Valor"], [("Páginas máximas", display_value(fmt.get("paginas_maximas"))), ("Fuente", display_value(fmt.get("fuente"))),
        ("Tamaño (pt)", display_value(fmt.get("tamano_pt"))), ("Interlineado", display_value(fmt.get("interlineado"))),
        ("Margen superior (cm)", display_value(fmt.get("margen_superior_cm"))), ("Margen inferior (cm)", display_value(fmt.get("margen_inferior_cm"))),
        ("Margen izquierdo (cm)", display_value(fmt.get("margen_izquierdo_cm"))), ("Margen derecho (cm)", display_value(fmt.get("margen_derecho_cm"))),
        ("Portada computa límite", boolean_value(memory.get("portada_computa_limite"))), ("Índice computa límite", boolean_value(memory.get("indice_computa_limite"))),
        ("Anexos computan límite", boolean_value(memory.get("anexos_computan_limite")))])
    doc.add_heading("Reglas, prohibiciones y consecuencias", 1)
    add_table_or_missing(doc, ["Tipo", "Descripción", "Consecuencia"], memory.get("reglas"),
        lambda item: (display_value(item.get("tipo")), display_value(item.get("descripcion")), display_value(item.get("consecuencia"))))


def _criteria(doc: Document, values: Any, title: str) -> None:
    doc.add_heading(title, 1)
    add_table_or_missing(doc, ["Criterio", "Puntuación máxima", "Tipo", "Dónde responder", "Descripción"], values,
        lambda item: (display_value(item.get("nombre")), display_value(item.get("puntuacion_maxima")), display_value(item.get("tipo")), display_value(item.get("donde_responder")), display_value(item.get("descripcion"))))


def _requirements(doc: Document, values: Any, title: str) -> None:
    doc.add_heading(title, 1)
    add_table_or_missing(doc, ["Descripción", "Tipo", "Obligatorio", "Dónde responder"], values,
        lambda item: (display_value(item.get("descripcion")), display_value(item.get("tipo")), boolean_value(item.get("obligatorio")), display_value(item.get("donde_responder"))))


def _pending(doc: Document, values: Any) -> None:
    doc.add_heading("Pendientes para preparar la oferta", 1)
    groups = (("documento_tecnico", "Documento técnico"), ("economico", "Información económica"), ("resumen", "Resumen y coordinación"), ("general", "Generales"))
    items = values if isinstance(values, list) else []
    for destination, title in groups:
        doc.add_heading(title, 2)
        rows = [(display_value(item.get("descripcion")), display_value(item.get("motivo")), display_value(item.get("destino"))) for item in items if isinstance(item, dict) and item.get("destino") == destination]
        add_table(doc, ["Pendiente", "Motivo", "Destino"], rows or [(MISSING, MISSING, destination)])


def _quality(doc: Document, value: Any) -> None:
    quality = value if isinstance(value, dict) else {}
    doc.add_heading("Calidad y limitaciones del análisis", 1)
    add_table(doc, ["Campo", "Valor"], [("Estado", display_value(quality.get("estado"))), ("Limitaciones", join_values(quality.get("limitaciones")))])


def _evidence(doc: Document, values: Any) -> None:
    doc.add_heading("Fuentes y evidencias", 1)
    add_table_or_missing(doc, ["Campo", "Documento", "Página", "Texto citado"], values,
        lambda item: (display_value(item.get("campo")), display_value(item.get("documento_id")), display_value(item.get("pagina")), truncate_value(item.get("texto_citado"))))
