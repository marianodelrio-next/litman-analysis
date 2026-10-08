"""Neutral, deterministic presentation helpers shared by DOCX renderers."""
from __future__ import annotations

from typing import Any, Callable

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

MISSING = "No consta"
BLUE, DARK, TEXT, MUTED, PALE = "1681BA", "123044", "1D2935", "667085", "EEF8FC"
AMBER, RED, GREEN = "B7791F", "B42318", "238636"


def boolean_value(value: Any) -> str:
    return "Sí" if value is True else "No" if value is False else MISSING


def display_value(value: Any) -> str:
    if value is None or value == "" or value == {} or value == []:
        return MISSING
    if isinstance(value, bool):
        return boolean_value(value)
    if isinstance(value, (list, tuple)):
        return join_values(value)
    if isinstance(value, dict):
        return named_description(value)
    return str(value)


def join_values(values: Any) -> str:
    return "; ".join(display_value(value) for value in values) if isinstance(values, list) and values else MISSING


def named_description(value: Any) -> str:
    if not isinstance(value, dict):
        return display_value(value)
    name, description = display_value(value.get("nombre")), display_value(value.get("descripcion"))
    return name if description == MISSING else f"{name} — {description}"


def amount_value(value: Any, currency: Any) -> str:
    return MISSING if value is None else f"{value} {currency}" if currency not in (None, "") else str(value)


def truncate_value(value: Any, limit: int = 500) -> str:
    return abbreviate_text(value, limit, " [cita abreviada]")


def abbreviate_text(value: Any, limit: int = 500, suffix: str = "") -> str:
    text = display_value(value)
    return text if len(text) <= limit else text[:limit - 1].rstrip() + "…" + suffix


def configure_document(doc: Document, *, title: str, subject: str, document_label: str, header_text: str) -> None:
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2), Cm(1.9)
    section.left_margin, section.right_margin = Cm(2), Cm(2)
    section.header_distance, section.footer_distance = Cm(.8), Cm(.8)
    doc.core_properties.title, doc.core_properties.subject = title, subject
    configure_styles(doc)
    header = section.header.paragraphs[0]; header.text = header_text; header.style = "LitMan Metadata"
    footer = section.footer.paragraphs[0]; footer.style = "LitMan Metadata"; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run(document_label + " · Página "); add_page_number(footer)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    for name, size, color in (("Normal", 11, TEXT), ("Title", 22, DARK), ("Heading 1", 22, DARK), ("Heading 2", 15, DARK), ("Heading 3", 12, BLUE)):
        style = styles[name]; style.font.name = "Arial"; style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size); style.font.color.rgb = RGBColor.from_string(color); style.paragraph_format.space_after = Pt(7); style.paragraph_format.line_spacing = 1.15
        if name.startswith("Heading"): style.paragraph_format.keep_with_next = True
    for name, size, color, bold in (("LitMan Lead", 12, TEXT, False), ("LitMan Fact Label", 8, MUTED, True), ("LitMan Fact Value", 10, TEXT, True), ("LitMan Callout", 10, DARK, False), ("LitMan Metadata", 8, MUTED, False), ("LitMan Card Title", 11, DARK, True), ("LitMan Card Body", 10, TEXT, False), ("LitMan Compact", 9, TEXT, False), ("LitMan Evidence", 9, MUTED, False)):
        if name not in styles: styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style = styles[name]; style.font.name = "Arial"; style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial"); style.font.size = Pt(size); style.font.bold = bold; style.font.color.rgb = RGBColor.from_string(color); style.paragraph_format.space_after = Pt(5); style.paragraph_format.line_spacing = 1.15


def add_page_number(paragraph: Any) -> None:
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); paragraph._p.append(field)


def set_cell_shading(cell: Any, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcpr.append(shd)


def set_repeat_table_header(row: Any) -> None:
    trpr = row._tr.get_or_add_trPr(); el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); trpr.append(el)


def prevent_row_split(row: Any) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def set_column_widths(table: Any, widths: list[float]) -> None:
    for row in table.rows:
        for cell, width in zip(row.cells, widths): cell.width = Cm(width)


def add_data_table(doc: Document, headers: list[str], rows: list[tuple[Any, ...]], widths: list[float] | None = None) -> Any:
    table = doc.add_table(rows=1, cols=len(headers)); table.autofit = False
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = str(value); set_cell_shading(cell, PALE)
        for p in cell.paragraphs: p.style = "LitMan Fact Label"
    set_repeat_table_header(table.rows[0]); prevent_row_split(table.rows[0])
    for values in rows:
        row = table.add_row(); prevent_row_split(row)
        for cell, value in zip(row.cells, values):
            cell.text = display_value(value)
            for p in cell.paragraphs: p.style = "LitMan Compact"
    if widths: set_column_widths(table, widths)
    return table


def add_fact_table(doc: Document, headers: list[str], values: list[Any]) -> Any:
    return add_data_table(doc, headers, [tuple(values)], [4.2] * len(headers))


def _paragraph_card(paragraph: Any, color: str, fill: str) -> None:
    """Add a non-tabular, accessible visual card treatment to a paragraph."""
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr"); left = OxmlElement("w:left")
    left.set(qn("w:val"), "single"); left.set(qn("w:sz"), "14"); left.set(qn("w:space"), "7"); left.set(qn("w:color"), color)
    borders.append(left); ppr.append(borders)
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); ppr.append(shd)
    paragraph.paragraph_format.left_indent = Cm(.18)
    paragraph.paragraph_format.space_before = Pt(5)


def add_callout(doc: Document, label: str, text: Any = None, *, kind: str = "neutral",
                show_body: bool = True) -> None:
    """Render a labelled card; its label remains meaningful without colour."""
    palettes = {"neutral": (BLUE, PALE), "warning": (AMBER, "FFF8E8"), "danger": (RED, "FDECEA"), "success": (GREEN, "EAF6EE")}
    color, fill = palettes.get(kind, palettes["neutral"])
    p = doc.add_paragraph(label, style="LitMan Card Title"); _paragraph_card(p, color, fill); set_keep_with_next(p)
    if not show_body:
        set_keep_together(p)
        return
    p = doc.add_paragraph(display_value(text), style="LitMan Callout"); _paragraph_card(p, color, fill); set_keep_together(p)


def add_empty_message(doc: Document, text: str) -> None: doc.add_paragraph(text, style="LitMan Callout")
def add_checkbox_item(doc: Document, text: Any) -> None: doc.add_paragraph("□ " + display_value(text), style="LitMan Card Body")
def add_label_value(doc: Document, label: str, value: Any, *, value_style: str = "LitMan Card Body") -> None:
    heading = doc.add_paragraph(label, style="LitMan Fact Label"); set_keep_with_next(heading)
    body = doc.add_paragraph(display_value(value), style=value_style); set_keep_together(body)
def set_keep_with_next(paragraph: Any) -> None: paragraph.paragraph_format.keep_with_next = True
def set_keep_together(paragraph: Any) -> None: paragraph.paragraph_format.keep_together = True


def add_table(doc: Document, headers: list[str], rows: list[tuple[Any, ...]]) -> None:
    add_data_table(doc, headers, rows)


def add_table_or_missing(doc: Document, headers: list[str], values: Any,
                         row: Callable[[dict[str, Any]], tuple[Any, ...]]) -> None:
    rows = [row(item) for item in values if isinstance(item, dict)] if isinstance(values, list) else []
    add_table(doc, headers, rows or [tuple(MISSING for _ in headers)])
