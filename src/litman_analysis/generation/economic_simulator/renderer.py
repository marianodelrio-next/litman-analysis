from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .planner import technical_max_points


NAVY = "002855"
BLUE = "4FA6DC"
ORANGE = "F2A900"
YELLOW = "FFF2CC"
GREY = "E7EAED"
GREEN = "008000"
RED = "C00000"
WHITE = "FFFFFF"


def create_clean_template(path: Path) -> None:
    workbook = Workbook()
    simulator = workbook.active
    simulator.title = "Simulador"
    criteria = workbook.create_sheet("Criterios y ayuda")
    checks = workbook.create_sheet("Comprobaciones")
    pending = workbook.create_sheet("Pendientes")
    neutral_plan = {
        "tender": {"name": "Licitación", "reference": None,
                   "budget_without_tax": None, "currency": "EUR"},
        "total_points": None,
        "price_formula": {"type": "no_localizada", "literal": None,
                          "max_points": None, "supported": False, "evidence": None},
        "criteria": [{"id": "criterion-001", "name": "Criterio configurable",
                      "type": "juicio_valor", "max_points": None,
                      "description": "Se sustituye con los criterios del pliego.",
                      "evidence": None}],
        "abnormally_low": {"aplica": None, "maximo_cualitativo": None,
                           "comparador": None, "reglas": []},
        "warnings": ["Plantilla neutra: generar desde un dossier validado."],
    }
    neutral_inputs = {
        "scenarios": [{
            "id": "scenario-01", "name": "Escenario configurable",
            "bidders": [
                {"name": "NUESTRA OFERTA", "price": None, "technical_score": None},
                {"name": "LICITADOR 1", "price": None, "technical_score": None},
            ],
        }],
        "notes": ["Completar al generar una licitación concreta."],
    }
    _render_simulator(simulator, neutral_plan, neutral_inputs)
    _render_criteria(criteria, neutral_plan)
    _render_checks(checks, neutral_plan, neutral_inputs)
    _render_pending(pending, neutral_plan, neutral_inputs)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.properties.title = "Plantilla de simulador económico interno"
    workbook.properties.subject = "Plantilla neutra; sin datos de licitación ni oferta"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def render_economic_workbook(template_path: Path, plan: dict, inputs: dict, output_path: Path) -> dict:
    workbook = load_workbook(template_path)
    for sheet in list(workbook.worksheets):
        workbook.remove(sheet)
    simulator = workbook.create_sheet("Simulador")
    criteria_sheet = workbook.create_sheet("Criterios y ayuda")
    checks = workbook.create_sheet("Comprobaciones")
    pending = workbook.create_sheet("Pendientes")
    _render_simulator(simulator, plan, inputs)
    _render_criteria(criteria_sheet, plan)
    _render_checks(checks, plan, inputs)
    _render_pending(pending, plan, inputs)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return validate_rendered_workbook(output_path, plan, inputs)


def validate_rendered_workbook(path: Path, plan: dict, inputs: dict) -> dict:
    errors, warnings = [], list(plan.get("warnings", []))
    try:
        workbook = load_workbook(path, data_only=False)
    except Exception as exc:
        return {"valid": False, "errors": [f"XLSX no válido: {exc}"], "warnings": warnings}
    expected = ["Simulador", "Criterios y ayuda", "Comprobaciones", "Pendientes"]
    if workbook.sheetnames != expected:
        errors.append("Las hojas del simulador no coinciden con la plantilla")
    if not plan.get("price_formula", {}).get("supported"):
        warnings.append("La puntuación económica requiere implementación o revisión manual.")
    technical_maximum = technical_max_points(plan.get("criteria", []))
    input_version = inputs.get("version")
    for scenario in inputs.get("scenarios", []):
        for bidder in scenario.get("bidders", []):
            label = f"{scenario.get('id', '?')}/{bidder.get('name', '?')}"
            has_aggregate = "technical_score" in bidder
            has_legacy = "criterion_scores" in bidder
            if has_aggregate and has_legacy:
                errors.append(f"{label}: mezcla ambigua de technical_score y criterion_scores")
                continue
            if input_version == "economic-inputs-v2" and not has_aggregate:
                errors.append(f"{label}: economic-inputs-v2 requiere technical_score")
                continue
            if input_version == "economic-inputs-v1" and not has_legacy:
                errors.append(f"{label}: economic-inputs-v1 requiere criterion_scores")
                continue
            if has_aggregate:
                value = bidder.get("technical_score")
                if value is not None and (not isinstance(value, (int, float)) or isinstance(value, bool)):
                    errors.append(f"{label}: technical_score debe ser numérica")
                elif value is not None and (value < 0 or value > technical_maximum):
                    errors.append(f"{label}: puntuación técnica fuera de rango (0–{technical_maximum})")
            elif has_legacy:
                criterion_ids = {item["id"] for item in plan.get("criteria", [])}
                unknown = set(bidder.get("criterion_scores", {})) - criterion_ids
                if unknown:
                    errors.append(f"{label}: criterios inexistentes: {sorted(unknown)}")
                for item in plan.get("criteria", []):
                    value = bidder.get("criterion_scores", {}).get(item["id"])
                    maximum = item.get("max_points")
                    if value is not None and (not isinstance(value, (int, float)) or isinstance(value, bool)
                                              or value < 0 or (maximum is not None and value > maximum)):
                        errors.append(f"{label}: puntuación fuera de rango en {item['name']}")
            else:
                errors.append(f"{label}: falta technical_score o criterion_scores")
    return {"valid": not errors, "errors": errors,
            "warnings": list(dict.fromkeys(warnings)), "recalculation": "on_open"}


def _render_simulator(sheet, plan: dict, inputs: dict) -> None:
    criteria = plan.get("criteria", [])
    technical_maximum = technical_max_points(criteria)
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells(start_row=1, start_column=1, end_row=2, end_column=10)
    sheet["A1"] = "SIMULADOR ECONÓMICO INTERNO"
    _title(sheet["A1"])
    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=10)
    tender_name = plan.get("tender", {}).get("name") or "Licitación"
    sheet["A3"] = (f"{tender_name} · Inputs amarillos · Importes sin IVA · "
                   "Las ofertas rivales son hipótesis internas")
    sheet["A3"].fill = PatternFill("solid", fgColor="DDEBF7")
    sheet["A3"].font = Font(italic=True, color=NAVY)

    _section(sheet, 5, "1. PARÁMETROS DEL PLIEGO", 10, NAVY)
    labels = [
        (6, "Expediente", plan.get("tender", {}).get("reference")),
        (7, "Presupuesto máximo (sin IVA)", plan.get("tender", {}).get("budget_without_tax")),
        (8, "Puntuación máxima precio", plan.get("price_formula", {}).get("max_points")),
        (9, "Puntuación total declarada", plan.get("total_points")),
    ]
    for row, label, value in labels:
        sheet.cell(row, 1, label)
        sheet.cell(row, 3, value)
        sheet.cell(row, 3).font = Font(color=GREEN)
        sheet.cell(row, 3).fill = PatternFill("solid", fgColor=GREY)
    sheet["E6"] = "Máximo no económico modelado"
    sheet["G6"] = technical_maximum
    sheet["E7"] = "Máximo cualitativo para umbral"
    sheet["G7"] = plan.get("abnormally_low", {}).get("maximo_cualitativo")
    sheet["E8"] = "Modelo de precio"
    sheet["G8"] = plan.get("price_formula", {}).get("type")
    sheet["E9"] = "Estado fórmula"
    sheet["G9"] = "SOPORTADA" if plan.get("price_formula", {}).get("supported") else "PENDIENTE"
    for coordinate in ("G6", "G7", "G8", "G9"):
        sheet[coordinate].font = Font(color=GREEN)
        sheet[coordinate].fill = PatternFill("solid", fgColor=GREY)

    current = 12
    for scenario_index, scenario in enumerate(inputs.get("scenarios", []), 1):
        end_column = 10
        _section(sheet, current, f"{scenario_index + 1}. ESCENARIO · {scenario['name'].upper()}", end_column,
                 ORANGE if scenario_index % 2 else BLUE)
        header_row = current + 2
        headers = ["Licitador", "Precio ofertado", "% bajada",
                   f"Puntuación técnica\n(0–{_display(technical_maximum)} p)",
                   "Puntos precio", "Puntuación total", "C cualitativo",
                   "Umbral baja anormal", "Estado", "Posición"]
        for column, header in enumerate(headers, 1):
            cell = sheet.cell(header_row, column, header)
            _header(cell)
        first_bidder = header_row + 1
        last_bidder = first_bidder + len(scenario.get("bidders", [])) - 1
        technical_col = 4
        price_col, total_col, coefficient_col, threshold_col, state_col, position_col = 5, 6, 7, 8, 9, 10
        price_range = f"$B${first_bidder}:$B${last_bidder}"
        total_range = f"${get_column_letter(total_col)}${first_bidder}:${get_column_letter(total_col)}${last_bidder}"
        for offset, bidder in enumerate(scenario.get("bidders", [])):
            row = first_bidder + offset
            sheet.cell(row, 1, bidder.get("name"))
            sheet.cell(row, 2, bidder.get("price"))
            _input(sheet.cell(row, 2))
            price_validation = DataValidation(
                type="decimal", operator="between", formula1="0.01",
                formula2="$C$7" if plan.get("tender", {}).get("budget_without_tax") is not None else "1E+15",
                allow_blank=True,
            )
            price_validation.error = "El precio debe ser positivo y no superar el presupuesto máximo."
            price_validation.errorTitle = "Precio no válido"
            sheet.add_data_validation(price_validation)
            price_validation.add(sheet.cell(row, 2))
            sheet.cell(row, 3, f'=IF(B{row}="","",1-B{row}/$C$7)')
            value = _bidder_technical_score(bidder, criteria)
            sheet.cell(row, technical_col, value)
            _input(sheet.cell(row, technical_col))
            if technical_maximum == 0:
                # Excel's decimal-between validation with equal bounds is unreliable.
                validation = DataValidation(type="list", formula1='"0"', allow_blank=True)
            else:
                validation = DataValidation(type="decimal", operator="between", formula1="0",
                                            formula2=str(technical_maximum), allow_blank=True)
            validation.error = "La puntuación técnica debe estar dentro del rango permitido."
            validation.errorTitle = "Puntuación técnica no válida"
            sheet.add_data_validation(validation)
            validation.add(sheet.cell(row, technical_col))
            price_letter = get_column_letter(price_col)
            if plan.get("price_formula", {}).get("supported"):
                formula = f'=IF(B{row}="","",IFERROR(MIN({price_range})/B{row}*$C$8,0))'
            else:
                formula = '="PENDIENTE"'
            sheet.cell(row, price_col, formula)
            if plan.get("price_formula", {}).get("supported"):
                sheet.cell(row, total_col,
                           f'=IF(OR(B{row}="",D{row}=""),"",SUM({price_letter}{row},D{row}))')
            else:
                sheet.cell(row, total_col, '=""')
            sheet.cell(row, coefficient_col,
                       f'=IF(OR(B{row}="",D{row}=""),"",IFERROR(D{row}/$G$7,0))')
            sheet.cell(row, threshold_col, _threshold_formula(plan, row, price_range, coefficient_col))
            comparator = "<" if plan.get("abnormally_low", {}).get("comparador") == "menor" else "<="
            sheet.cell(row, state_col,
                       f'=IF(B{row}="","",IF({get_column_letter(threshold_col)}{row}="","NO CALCULABLE",IF(B{row}{comparator}{get_column_letter(threshold_col)}{row},"POSIBLE BAJA ANORMAL","NORMAL")))')
            sheet.cell(row, position_col, f'=IF({get_column_letter(total_col)}{row}="","",RANK({get_column_letter(total_col)}{row},{total_range},0))')
        _scenario_formats(sheet, first_bidder, last_bidder, len(headers), state_col)
        current = last_bidder + 4
    sheet.freeze_panes = "A12"
    sheet.column_dimensions["A"].width = 22
    for column in range(2, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 16
    sheet.auto_filter.ref = None


def _threshold_formula(plan: dict, row: int, price_range: str, coefficient_col: int) -> str:
    abnormal = plan.get("abnormally_low", {})
    if (abnormal.get("aplica") is not True or not abnormal.get("reglas")
            or not abnormal.get("maximo_cualitativo")):
        return '=""'
    coefficient = f"{get_column_letter(coefficient_col)}{row}"
    fallback = '""'
    for rule in reversed(abnormal.get("reglas", [])):
        minimum = rule.get("min_licitadores", 1)
        maximum = rule.get("max_licitadores")
        condition = f"COUNT({price_range})>={minimum}"
        if maximum is not None:
            condition = f"AND({condition},COUNT({price_range})<={maximum})"
        if rule.get("base") is None or rule.get("referencia") == "otra":
            expression = '""'
        else:
            reference = f"AVERAGE({price_range})" if rule.get("referencia") == "media_ofertas" else "$C$7"
            expression = f"({rule['base']}+{coefficient}*{rule.get('peso_calidad') or 0})*{reference}"
        fallback = f"IF({condition},{expression},{fallback})"
    return f'=IF(OR(B{row}="",{coefficient}=""),"",IFERROR({fallback},""))'


def _bidder_technical_score(bidder: dict, criteria: list[dict]):
    """Use v2 aggregate input, or derive it for legacy v1 rerendering only."""
    if "technical_score" in bidder and "criterion_scores" in bidder:
        raise ValueError("Un licitador no puede mezclar technical_score y criterion_scores.")
    if "technical_score" in bidder:
        return bidder.get("technical_score")
    if "criterion_scores" in bidder:
        scores = bidder.get("criterion_scores") or {}
        values = [scores.get(item["id"]) for item in criteria]
        if not any(value is not None for value in values):
            return None
        return sum(float(value or 0) for value in values)
    return None


def _render_criteria(sheet, plan: dict) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:G2")
    sheet["A1"] = "CRITERIOS, FÓRMULAS Y TRAZABILIDAD"
    _title(sheet["A1"])
    headers = ["Criterio", "Máximo", "Tipo", "Tratamiento", "Documento", "Página", "Texto citado"]
    for column, header in enumerate(headers, 1):
        _header(sheet.cell(4, column, header))
    rows = [{"name": "Oferta económica", "max_points": plan.get("price_formula", {}).get("max_points"),
             "type": "precio", "description": plan.get("price_formula", {}).get("literal"),
             "evidence": plan.get("price_formula", {}).get("evidence")}] + [
        {"name": item["name"], "max_points": item.get("max_points"), "type": item.get("type"),
         "description": item.get("description"), "evidence": item.get("evidence")}
        for item in plan.get("criteria", [])
    ]
    for row_index, item in enumerate(rows, 5):
        evidence = item.get("evidence") or {}
        values = [item["name"], item.get("max_points"), item.get("type"), item.get("description"),
                  evidence.get("documento_id"), evidence.get("pagina"), evidence.get("texto_citado")]
        for column, value in enumerate(values, 1):
            sheet.cell(row_index, column, value)
    warning_row = 7 + len(rows)
    sheet.cell(warning_row, 1, "ADVERTENCIAS")
    sheet.cell(warning_row, 1).font = Font(bold=True, color=RED)
    for offset, warning in enumerate(plan.get("warnings", []), 1):
        sheet.cell(warning_row + offset, 1, warning)
        sheet.merge_cells(start_row=warning_row + offset, start_column=1, end_row=warning_row + offset, end_column=7)
    sheet.freeze_panes = "A5"
    widths = [30, 12, 24, 42, 22, 10, 55]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _render_checks(sheet, plan: dict, inputs: dict) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:F2")
    sheet["A1"] = "COMPROBACIONES DEL MODELO"
    _title(sheet["A1"])
    headers = ["Control", "Actual", "Esperado", "Diferencia", "Estado", "Observaciones"]
    for column, header in enumerate(headers, 1):
        _header(sheet.cell(4, column, header))
    qualitative = technical_max_points(plan.get("criteria", []))
    modeled = qualitative + (plan.get("price_formula", {}).get("max_points") or 0)
    checks = [
        ("Total máximo modelado", modeled, plan.get("total_points"), "La suma debe coincidir con el pliego."),
        ("Fórmula económica soportada", 1 if plan.get("price_formula", {}).get("supported") else 0, 1,
         "Una fórmula no soportada requiere revisión manual."),
        ("Máximo no económico modelado", qualitative, qualitative, "Suma de criterios no económicos."),
    ]
    for row, (label, actual, expected, note) in enumerate(checks, 5):
        sheet.cell(row, 1, label)
        sheet.cell(row, 2, actual)
        sheet.cell(row, 3, expected)
        sheet.cell(row, 4, f'=IF(OR(B{row}="",C{row}=""),"",B{row}-C{row})')
        sheet.cell(row, 5, f'=IF(D{row}="","ADVERTENCIA",IF(ABS(D{row})<0.000001,"OK","ERROR"))')
        sheet.cell(row, 6, note)
    sheet.column_dimensions["A"].width = 35
    sheet.column_dimensions["F"].width = 55
    for column in ("B", "C", "D", "E"):
        sheet.column_dimensions[column].width = 16


def _render_pending(sheet, plan: dict, inputs: dict) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:C2")
    sheet["A1"] = "PENDIENTES DEL SIMULADOR"
    _title(sheet["A1"])
    for column, header in enumerate(("Pendiente", "Origen", "Acción"), 1):
        _header(sheet.cell(4, column, header))
    items = []
    if not plan.get("price_formula", {}).get("supported"):
        items.append(("Fórmula económica no soportada", "Pliego", "Revisar e implementar antes de simular"))
    items.extend((warning, "Validación", "Revisar") for warning in plan.get("warnings", []))
    items.extend((note, "Inputs internos", "Completar") for note in inputs.get("notes", []))
    for row, values in enumerate(items, 5):
        for column, value in enumerate(values, 1):
            sheet.cell(row, column, value)
    sheet.column_dimensions["A"].width = 65
    sheet.column_dimensions["B"].width = 22
    sheet.column_dimensions["C"].width = 42


def _scenario_formats(sheet, first_row: int, last_row: int, columns: int, state_col: int) -> None:
    thin = Side(style="thin", color="B7B7B7")
    for row in sheet.iter_rows(min_row=first_row, max_row=last_row, min_col=1, max_col=columns):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    state_range = f"{get_column_letter(state_col)}{first_row}:{get_column_letter(state_col)}{last_row}"
    sheet.conditional_formatting.add(state_range,
                                     FormulaRule(formula=[f'{get_column_letter(state_col)}{first_row}="POSIBLE BAJA ANORMAL"'],
                                                 font=Font(color=RED, bold=True)))
    for row in range(first_row, last_row + 1):
        sheet.cell(row, 2).number_format = '#,##0.00 [$€-es-ES]'
        sheet.cell(row, 3).number_format = '0.00%'


def _title(cell) -> None:
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(color=WHITE, bold=True, size=18)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _section(sheet, row: int, text: str, end_column: int, color: str) -> None:
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end_column)
    cell = sheet.cell(row, 1, text)
    cell.fill = PatternFill("solid", fgColor=color)
    cell.font = Font(color=WHITE, bold=True, size=12)


def _header(cell) -> None:
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(color=WHITE, bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _input(cell) -> None:
    cell.fill = PatternFill("solid", fgColor=YELLOW)
    cell.font = Font(color="0070C0")


def _display(value) -> str:
    if value is None:
        return "?"
    return str(int(value)) if float(value).is_integer() else str(value)
