from __future__ import annotations
from pathlib import Path

from docx import Document
def write_summary(data: dict, path: Path) -> None:
    doc, head = Document(), data["cabecera"]
    doc.add_heading(head.get("nombre") or "Resumen de licitación", 0)
    doc.add_paragraph(" | ".join(f"{label}: {value}" for label, value in
        (("Expediente", head.get("expediente")), ("Organismo", head.get("organismo")),
         ("Fecha límite", head.get("fecha_limite")), ("Presupuesto sin IVA", head.get("presupuesto_sin_iva"))) if value is not None))
    sections = [("Resumen técnico", [data.get("contexto")]), ("Objetivos", data.get("objetivos", [])),
        ("Alcance", data.get("alcance", [])), ("Fases", [_phase(item) for item in data.get("fases", [])]),
        ("Equipo exigido", [_entity(item) for item in data.get("equipo", [])]), ("Tecnologías", data.get("tecnologias", [])),
        ("Documentos que presentar", [_offer(item) for item in data.get("documentos_oferta", [])]),
        ("Criterios de adjudicación", [_criterion(item) for item in data.get("criterios_adjudicacion", [])]),
        ("Requisitos concretos", [_requirement(item) for item in data.get("requisitos_oferta", [])]),
        ("Pendientes", [_pending(item) for item in data.get("pendientes", [])])]
    for title, entries in sections:
        doc.add_heading(title, level=1)
        clean = [item for item in entries if item]
        if not clean: doc.add_paragraph("[PENDIENTE: completar]")
        for item in clean: doc.add_paragraph(str(item), style="List Bullet")
    doc.save(path)


def _phase(item): return " — ".join(part for part in [item.get("nombre"), item.get("descripcion"), item.get("plazo")] if part)
def _entity(item): return " — ".join(part for part in [item.get("nombre"), item.get("descripcion")] if part)
def _requirement(item): return item.get("descripcion", "") + (f" ({item['donde_responder']})" if item.get("donde_responder") else "")
def _offer(item): return item.get("nombre", "") + (f" — {item['sobre']}" if item.get("sobre") else "") + (f": {item['descripcion']}" if item.get("descripcion") else "")
def _criterion(item): return item.get("nombre", "") + (f" ({item['puntuacion_maxima']} puntos)" if item.get("puntuacion_maxima") is not None else "") + (f" — responder en {item['donde_responder']}" if item.get("donde_responder") else "") + (f": {item['descripcion']}" if item.get("descripcion") else "")
def _pending(item): return f"[PENDIENTE] {item.get('descripcion')}" + (f" — {item['motivo']}" if item.get("motivo") else "")
