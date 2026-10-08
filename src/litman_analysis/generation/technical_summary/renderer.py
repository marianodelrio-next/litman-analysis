from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document

from ..docx_helpers import (MISSING, add_table, add_table_or_missing, amount_value,
                            boolean_value, display_value, join_values, named_description)


def write_technical_summary(data: dict[str, Any], path: Path) -> None:
    """Render the dossier-only technical view (never proposal content)."""
    doc = Document()
    _title_and_header(doc, "Resumen técnico", data.get("cabecera") or {})
    _text_section(doc, "Contexto y situación de partida", data.get("contexto"))
    _list_section(doc, "Objetivos", data.get("objetivos"))
    _list_section(doc, "Alcance", data.get("alcance"))
    _list_section(doc, "Trabajos requeridos", data.get("trabajos"))
    _phases(doc, data.get("fases"))
    _entities(doc, "Entregables del contrato", data.get("entregables_contrato"))
    _entities(doc, "Equipo y perfiles exigidos", data.get("equipo"))
    _list_section(doc, "Tecnologías", data.get("tecnologias"))
    _requirements(doc, data.get("requisitos_oferta"), "Requisitos que condicionan la solución técnica")
    _criteria(doc, data.get("criterios_adjudicacion"), "Criterios que condicionan la propuesta")
    _quality(doc, data.get("calidad"))
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def _title_and_header(doc: Document, title: str, header: dict[str, Any]) -> None:
    doc.add_heading(title, 0)
    add_table(doc, ["Campo", "Valor"], [
        ("Oportunidad", display_value(header.get("nombre"))), ("Expediente", display_value(header.get("expediente"))),
        ("Organismo", display_value(header.get("organismo"))), ("Fecha límite", display_value(header.get("fecha_limite"))),
        ("Presupuesto sin IVA", amount_value(header.get("presupuesto_sin_iva"), header.get("moneda"))),
        ("Moneda", display_value(header.get("moneda"))),
    ])


def _text_section(doc: Document, title: str, value: Any) -> None:
    doc.add_heading(title, 1)
    doc.add_paragraph(display_value(value))


def _list_section(doc: Document, title: str, values: Any) -> None:
    doc.add_heading(title, 1)
    if not isinstance(values, list) or not values:
        doc.add_paragraph(MISSING)
        return
    for value in values:
        doc.add_paragraph(display_value(value), style="List Bullet")


def _phases(doc: Document, phases: Any) -> None:
    doc.add_heading("Fases, plazos y entregables", 1)
    if not isinstance(phases, list) or not phases:
        doc.add_paragraph(MISSING)
        return
    for index, item in enumerate(phases, 1):
        item = item if isinstance(item, dict) else {}
        doc.add_heading(f"{index}. {display_value(item.get('nombre'))}", 2)
        add_table(doc, ["Descripción", "Plazo"], [(display_value(item.get("descripcion")), display_value(item.get("plazo")))])
        doc.add_paragraph("Entregables:")
        values = item.get("entregables")
        if isinstance(values, list) and values:
            for value in values:
                doc.add_paragraph(named_description(value), style="List Bullet")
        else:
            doc.add_paragraph(MISSING, style="List Bullet")


def _entities(doc: Document, title: str, values: Any) -> None:
    doc.add_heading(title, 1)
    add_table_or_missing(doc, ["Nombre", "Descripción"], values,
                         lambda item: (display_value(item.get("nombre")), display_value(item.get("descripcion"))))


def _criteria(doc: Document, values: Any, title: str) -> None:
    doc.add_heading(title, 1)
    add_table_or_missing(doc, ["Criterio", "Puntuación máxima", "Tipo", "Dónde responder", "Descripción"], values,
        lambda item: (display_value(item.get("nombre")), display_value(item.get("puntuacion_maxima")), display_value(item.get("tipo")), display_value(item.get("donde_responder")), display_value(item.get("descripcion"))))


def _requirements(doc: Document, values: Any, title: str) -> None:
    doc.add_heading(title, 1)
    add_table_or_missing(doc, ["Descripción", "Tipo", "Obligatorio", "Dónde responder"], values,
        lambda item: (display_value(item.get("descripcion")), display_value(item.get("tipo")), boolean_value(item.get("obligatorio")), display_value(item.get("donde_responder"))))


def _quality(doc: Document, value: Any) -> None:
    quality = value if isinstance(value, dict) else {}
    doc.add_heading("Calidad y limitaciones del análisis", 1)
    add_table(doc, ["Campo", "Valor"], [("Estado", display_value(quality.get("estado"))), ("Limitaciones", join_values(quality.get("limitaciones")))])
