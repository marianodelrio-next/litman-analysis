from __future__ import annotations

import json
import re
import sys
import unicodedata
from pathlib import Path
from xml.etree import ElementTree

from litman_analysis.domain.documents import Document, NormalizedDocument


def normalize_documents(documents: list[Document], output_dir: Path,
                        ) -> list[NormalizedDocument]:
    output_dir.mkdir(parents=True, exist_ok=True)
    results = []
    for index, doc in enumerate(documents, 1):
        print(f"[normalización {index}/{len(documents)}] {doc.filename}", file=sys.stderr, flush=True)
        result = normalize_document(doc, output_dir)
        print(f"  → {result.status} ({result.method or 'sin extracción'})", file=sys.stderr, flush=True)
        results.append(result)
    return results


def normalize_document(document: Document, output_dir: Path,
                       ) -> NormalizedDocument:
    if document.status != "ready":
        return NormalizedDocument(document.document_id, None, document.status, None, None, 0,
                                  [document.status_reason or "Documento no procesable"])
    try:
        content, method, pages, warnings = _to_markdown(document)
    except Exception as exc:  # el expediente debe poder continuar con otros documentos
        return NormalizedDocument(document.document_id, None, "error", None, None, 0,
                                  [f"Error de extracción: {exc}"])
    if not content.strip():
        return NormalizedDocument(document.document_id, None, "empty", method, pages, 0,
                                  ["No se obtuvo texto"])
    path = output_dir / f"{document.document_id}.md"
    path.write_text(content, encoding="utf-8")
    return NormalizedDocument(document.document_id, str(path), "ready", method, pages,
                              len(content), warnings)


def _to_markdown(document: Document) -> tuple[str, str, int | None, list[str]]:
    suffix = document.suffix
    header = f"<!-- documento_id: {document.document_id} | archivo: {document.filename} -->\n\n"
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            return "", "pdf_unavailable", None, ["Instala pypdf para extraer texto del PDF"]
        reader = PdfReader(str(document.path))
        pages = []
        for number, page in enumerate(reader.pages, 1):
            text = _clean_pdf_text(page.extract_text() or "")
            pages.append(f"<!-- pagina: {number} -->\n## Página {number}\n\n{text.strip()}\n")
        content = header + "\n".join(pages)
        warnings = ["El PDF puede contener tablas o elementos visuales no representados"]
        if not content.replace(header, "").strip():
            warnings.append("PDF sin texto extraíble; requiere OCR")
        return content, "pdf_text_by_page", len(reader.pages), warnings
    if suffix == ".docx":
        from docx import Document as DocxDocument
        source = DocxDocument(document.path)
        blocks = [p.text.strip() for p in source.paragraphs if p.text.strip()]
        for table in source.tables:
            rows = [[cell.text.strip().replace("|", "\\|") for cell in row.cells] for row in table.rows]
            if rows:
                blocks.append(_table_markdown(rows))
        return header + "\n\n".join(blocks), "docx_paragraphs_tables", None, []
    raw = document.path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".json":
        try:
            raw = json.dumps(json.loads(raw), ensure_ascii=False, indent=2)
        except json.JSONDecodeError:
            pass
    elif suffix == ".xml":
        try:
            root = ElementTree.fromstring(raw)
            raw = "\n".join((element.text or "").strip() for element in root.iter() if (element.text or "").strip())
        except ElementTree.ParseError:
            pass
    return header + raw, f"{suffix[1:]}_text", None, []


def _clean_pdf_text(value: str) -> str:
    """Remove unambiguous PDF encoding noise without guessing word boundaries."""
    value = unicodedata.normalize("NFKC", value).replace("\u00ad", "")
    return re.sub(r"(?<=\w)-[ \t]*\r?\n[ \t]*(?=\w)", "-", value)


def _table_markdown(rows: list[list[str]]) -> str:
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    return "\n".join(["| " + " | ".join(rows[0]) + " |",
                       "| " + " | ".join("---" for _ in range(width)) + " |"] +
                      ["| " + " | ".join(row) + " |" for row in rows[1:]])
