from __future__ import annotations
from pathlib import Path
from typing import Any
from docx import Document
from ..docx_helpers import MISSING, add_callout, add_checkbox_item, add_data_table, add_fact_table, add_label_value, amount_value, boolean_value, configure_document, display_value, named_description

GROUPS = (('documento_tecnico', 'Documento técnico'), ('economico', 'Información económica'), ('resumen', 'Resumen y coordinación'), ('general', 'Generales'))
def _items(x: Any) -> list[Any]: return x if isinstance(x, list) else []
def _memory(d: dict[str, Any]) -> dict[str, Any]: return d.get('memoria_tecnica') if isinstance(d.get('memoria_tecnica'), dict) else {}
def _rules(d: dict[str, Any]) -> list[Any]: return _items(_memory(d).get('reglas'))
def _header(h: dict[str, Any]) -> str: return ' · '.join(str(x) for x in (h.get('expediente'), h.get('nombre')) if x not in (None, '')) or MISSING

def write_bid_preparation_guide(data: dict[str, Any], path: Path) -> None:
    doc = Document(); h = data.get('cabecera') if isinstance(data.get('cabecera'), dict) else {}
    configure_document(doc, title='Guía de preparación de la oferta', subject='Guía operativa de preparación de oferta', document_label='Guía de preparación de la oferta', header_text=_header(h))
    _intro(doc, data, h); _context(doc, data); _documents(doc, data); _requirements(doc, data); _criteria(doc, data); _pending(doc, data)
    path.parent.mkdir(parents=True, exist_ok=True); doc.save(path)

def _intro(doc: Document, d: dict[str, Any], h: dict[str, Any]) -> None:
    doc.add_heading('GUÍA DE PREPARACIÓN DE LA OFERTA', 1); doc.add_paragraph(display_value(h.get('nombre')), style='LitMan Lead')
    meta = ' · '.join(x for x in (display_value(h.get('organismo')) if h.get('organismo') else '', f"Expediente {h['expediente']}" if h.get('expediente') else '') if x)
    doc.add_paragraph(meta or MISSING, style='LitMan Metadata'); docs = _items(d.get('documentos_oferta')); pending = _items(d.get('pendientes'))
    add_fact_table(doc, ['FECHA LÍMITE', 'PRESUPUESTO', 'DOCUMENTOS OBLIGATORIOS', 'PENDIENTES'], [h.get('fecha_limite'), amount_value(h.get('presupuesto_sin_iva'), h.get('moneda')), sum(1 for x in docs if isinstance(x, dict) and x.get('obligatorio') is True), len(pending)])
    critical = []
    for r in _rules(d):
        if isinstance(r, dict):
            kind, consequence = str(r.get('tipo', '')).lower(), display_value(r.get('consecuencia'))
            if kind == 'prohibicion': critical.append(('PROHIBICIÓN', r.get('descripcion'), 'danger'))
            if 'exclu' in consequence.lower(): critical.append(('POSIBLE EXCLUSIÓN', consequence, 'danger'))
    for label, text, kind in critical: add_callout(doc, label, text, kind=kind)

def _context(doc: Document, d: dict[str, Any]) -> None:
    doc.add_heading('1. Contexto', 2); doc.add_paragraph(display_value(d.get('contexto')), style='LitMan Callout')

def _documents(doc: Document, d: dict[str, Any]) -> None:
    doc.add_heading('2. Documentación que presentar', 2); values = [x for x in _items(d.get('documentos_oferta')) if isinstance(x, dict)]
    rows = [(x.get('nombre'), x.get('sobre'), boolean_value(x.get('obligatorio')), x.get('descripcion')) for x in values]
    add_data_table(doc, ['DOCUMENTO', 'SOBRE', 'OBLIGATORIEDAD', 'OBSERVACIÓN'], rows or [(MISSING,) * 4], [5, 2, 3, 6])
    doc.add_paragraph('Estos documentos corresponden a la presentación de la oferta, no a los entregables posteriores del contrato.', style='LitMan Callout')

def _requirements(doc: Document, d: dict[str, Any]) -> None:
    doc.add_heading('3. Requisitos para preparar la oferta', 2); m = _memory(d)
    for x in (x for x in _items(d.get('requisitos_oferta')) if isinstance(x, dict)):
        label = 'OBLIGATORIO' if x.get('obligatorio') is True else 'NO OBLIGATORIO' if x.get('obligatorio') is False else 'OBLIGATORIEDAD NO INDICADA'
        add_callout(doc, label + ' · ' + display_value(x.get('tipo')), x.get('descripcion')); add_label_value(doc, 'Dónde responder', x.get('donde_responder'))
    sections = [x for x in _items(m.get('secciones')) if isinstance(x, dict)]
    if sections:
        doc.add_heading('Estructura exigida de la memoria', 3)
        for x in sorted(sections, key=lambda y: (y.get('orden') is None, y.get('orden', 0))):
            title = f"{x.get('orden'):02d}" if isinstance(x.get('orden'), int) else MISSING
            add_callout(doc, title + ' · ' + display_value(x.get('titulo')), x.get('descripcion')); add_label_value(doc, 'Obligatoriedad', boolean_value(x.get('obligatoria'))); add_label_value(doc, 'Debe responder a', x.get('requisitos_relacionados'))
    else:
        suggested = _items(d.get('estructura_documento_tecnico'))
        if suggested:
            doc.add_heading('Estructura sugerida de la memoria', 3)
            for x in suggested: add_checkbox_item(doc, named_description(x))
    for label, key in (('Memoria técnica requerida', 'requerida'), ('Nombre del documento', 'nombre_documento'), ('Idioma', 'idioma'), ('Estructura impuesta', 'estructura_impuesta')): add_label_value(doc, label, boolean_value(m.get(key)) if key in ('requerida', 'estructura_impuesta') else m.get(key))
    f = m.get('formato') if isinstance(m.get('formato'), dict) else {}; doc.add_heading('Formato y límites', 3)
    fields = (('Máximo de páginas', 'paginas_maximas', ' páginas'), ('Fuente', 'fuente', ''), ('Tamaño', 'tamano_pt', ' pt'), ('Interlineado', 'interlineado', ''), ('Margen superior', 'margen_superior_cm', ' cm'), ('Margen inferior', 'margen_inferior_cm', ' cm'), ('Margen izquierdo', 'margen_izquierdo_cm', ' cm'), ('Margen derecho', 'margen_derecho_cm', ' cm'))
    for label, key, unit in fields:
        if f.get(key) is not None: add_checkbox_item(doc, f'{label}: {f[key]}{unit}')
    missing = _missing(f, m)
    if missing: add_callout(doc, 'REVISAR', 'No consta: ' + ', '.join(missing), kind='warning')
    for label, key, verb in (('La portada', 'portada_computa_limite', 'computa'), ('El índice', 'indice_computa_limite', 'computa'), ('Los anexos', 'anexos_computan_limite', 'computan')):
        if m.get(key) is not None: add_checkbox_item(doc, f'{label} {verb if m[key] else "no " + verb} en el límite.')
    for x in _rules(d):
        if isinstance(x, dict):
            kind = str(x.get('tipo', '')).lower(); add_callout(doc, display_value(x.get('tipo')), x.get('descripcion'), kind='danger' if kind == 'prohibicion' else 'warning' if kind == 'advertencia' else 'neutral'); add_label_value(doc, 'Consecuencia', x.get('consecuencia'))

def _missing(f: dict[str, Any], m: dict[str, Any]) -> list[str]:
    missing = [label for key, label in (('fuente', 'fuente'), ('tamano_pt', 'tamaño'), ('interlineado', 'interlineado'), ('margen_superior_cm', 'margen superior'), ('margen_inferior_cm', 'margen inferior'), ('margen_izquierdo_cm', 'margen izquierdo'), ('margen_derecho_cm', 'margen derecho'), ('paginas_maximas', 'límite máximo de páginas')) if f.get(key) is None]
    missing.extend(label for key, label in (('portada_computa_limite', 'si la portada computa en el límite'), ('indice_computa_limite', 'si el índice computa en el límite'), ('anexos_computan_limite', 'si los anexos computan en el límite')) if m.get(key) is None)
    return missing

def _criteria(doc: Document, d: dict[str, Any]) -> None:
    doc.add_heading('4. Criterios de adjudicación', 2); values = [x for x in _items(d.get('criterios_adjudicacion')) if isinstance(x, dict)]
    rows = [(x.get('nombre'), x.get('puntuacion_maxima'), x.get('tipo'), x.get('donde_responder')) for x in values]; add_data_table(doc, ['CRITERIO', 'PUNTUACIÓN', 'TIPO', 'DÓNDE RESPONDER'], rows or [(MISSING,) * 4], [5, 3, 3, 5])

def _pending(doc: Document, d: dict[str, Any]) -> None:
    doc.add_heading('5. Pendientes y comprobación final', 2); values = [x for x in _items(d.get('pendientes')) if isinstance(x, dict)]
    for code, title in GROUPS:
        group = [x for x in values if x.get('destino') == code]
        if group:
            doc.add_heading(title, 3)
            for x in group: add_checkbox_item(doc, x.get('descripcion')); add_label_value(doc, 'Motivo', x.get('motivo'))
    q = d.get('calidad') if isinstance(d.get('calidad'), dict) else {}; limits = _items(q.get('limitaciones'))
    if limits: add_callout(doc, 'LIMITACIONES RELEVANTES', '; '.join(display_value(x) for x in limits), kind='warning')
