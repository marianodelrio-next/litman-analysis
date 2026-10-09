from __future__ import annotations

import re
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


HISTORICAL_TERMS = ("OCS", "DDnA", "Olympic Data")

UI_TEXT = {"cover_label": "Propuesta técnica", "index_title": "Índice", "guidance": "Guía de redacción", "labels": ["Finalidad", "Qué espera encontrar el evaluador", "Preguntas que debe responder", "Requisitos obligatorios", "Criterios de valoración", "Información disponible en el expediente", "Información que debe aportar la empresa", "Evidencias", "Evitar", "Puntuación máxima", "Extensión máxima", "Extensión orientativa"]}

DEFAULT_HEADING_LEVELS = {
    1: {"font": "Roboto", "font_size_pt": 16, "bold": True, "color": "4FA6DC"},
    2: {"font": "M PLUS Rounded 1c Medium", "font_size_pt": 13, "bold": False, "color": "F2A900"},
    3: {"font": "M PLUS Rounded 1c Medium", "font_size_pt": 13, "bold": False, "color": "A36785"},
}

DEFAULT_TOC_LEVELS = {
    1: {"font": "Roboto Light", "font_size_pt": 11, "bold": False, "italic": False, "color": "000000"},
    2: {"font": "Roboto Light", "font_size_pt": 11, "bold": False, "italic": False, "color": "000000"},
}


def render_technical_document(template_path: Path, content: dict, plan: dict, output_path: Path,
                              *, include_internal_pending: bool = True,
                              include_editorial_guidance: bool = True,
                              body_marker: str = "[[DOCUMENT_BODY]]",
                              table_format: dict | None = None,
                              toc_format: dict | None = None,
                              heading_format: dict | None = None) -> dict:
    document = Document(template_path)
    _replace_interface_text(document, content["document_title"], plan)
    _apply_format(document, plan.get("formal_constraints", {}))
    marker = next((paragraph for paragraph in document.paragraphs if paragraph.text.strip() == body_marker), None)
    if marker is None:
        raise ValueError(f"La plantilla no contiene el marcador de cuerpo {body_marker}")

    toc_entries = []
    planned_sections = {item["id"]: item for item in plan.get("sections", [])}
    for index, section in enumerate(content["sections"], 1):
        planned = planned_sections.get(section["id"], section)
        visible_title = _display_title({**planned, "title": section["title"]})
        heading = _insert_paragraph_before(marker, visible_title, "Heading 1")
        _apply_heading_format(document, heading, 1, heading_format or {},
                              plan.get("formal_constraints", {}))
        _add_bookmark(heading, _bookmark_name(section["id"]), index)
        toc_entries.append({"id": section["id"], "title": f"{index}. {visible_title}", "level": 1})
        if include_editorial_guidance:
            _insert_editorial_guidance(document, marker, planned.get("guidance", {}), plan.get("language", "es"))
        heading_counts = {level: 0 for level in range(2, 7)}
        for block in section["blocks"]:
            if block["type"] == "subheading":
                level = block.get("level", 2)
                heading_counts[level] += 1
                for deeper_level in range(level + 1, 7):
                    heading_counts[deeper_level] = 0
                for parent_level in range(2, level):
                    if heading_counts[parent_level] == 0:
                        heading_counts[parent_level] = 1
                label = ".".join(str(value) for value in
                                 [index] + [heading_counts[item] for item in range(2, level + 1)])
                subheading = _insert_paragraph_before(marker, block["text"], f"Heading {level}")
                _apply_heading_format(document, subheading, level, heading_format or {},
                                      plan.get("formal_constraints", {}))
                entry_id = block.get("id") or f"{section['id']}-subsection-{level}-{heading_counts[level]}"
                bookmark_id = index * 10000 + level * 1000 + heading_counts[level]
                _add_bookmark(subheading, _bookmark_name(entry_id), bookmark_id)
                toc_entries.append({"id": entry_id, "title": f"{label}. {block['text']}", "level": level})
            else:
                _insert_block(document, marker, block, include_internal_pending, table_format or {})
        if include_internal_pending:
            for pending in section.get("pending", []):
                _insert_paragraph_before(marker, f"[PENDIENTE] {pending}", "LitMan Analysis Pending")
    _populate_toc(document, toc_entries, toc_format or {})
    marker._element.getparent().remove(marker._element)
    _mark_fields_for_update(document)
    document.save(output_path)
    return validate_rendered_document(output_path, content, plan)


def validate_rendered_document(path: Path, content: dict, plan: dict) -> dict:
    errors = []
    warnings = []
    try:
        with zipfile.ZipFile(path) as archive:
            broken = archive.testzip()
            if broken:
                errors.append(f"Parte DOCX corrupta: {broken}")
    except zipfile.BadZipFile:
        errors.append("El resultado no es un DOCX válido")
        return {"valid": False, "errors": errors, "warnings": warnings}

    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    if "[[DOCUMENT_BODY]]" in text:
        errors.append("Queda el marcador interno del cuerpo")
    for section in content["sections"]:
        if section["title"] not in text:
            errors.append(f"Falta la sección: {section['title']}")
    found_historical = [term for term in HISTORICAL_TERMS if re.search(rf"\b{re.escape(term)}\b", text, re.I)]
    if found_historical:
        warnings.append(
            "Revisar posibles referencias históricas (pueden proceder legítimamente del expediente): "
            + ", ".join(found_historical)
        )
    max_pages = plan.get("formal_constraints", {}).get("paginas_maximas")
    if max_pages:
        warnings.append(
            f"El límite es {max_pages} páginas, pero requiere validación mediante renderizado PDF; "
            "no puede comprobarse de forma fiable leyendo el DOCX."
        )
    return {"valid": not errors, "errors": errors, "warnings": warnings,
            "page_limit_validation": "not_performed" if max_pages else "not_required"}


def _replace_interface_text(document: Document, title: str, plan: dict) -> None:
    _replace_marker(document, "[[COVER_LABEL]]", UI_TEXT["cover_label"])
    _replace_marker(document, "[[COVER_TITLE]]", title)
    index_title = _replace_marker(document, "[[INDEX_TITLE]]", UI_TEXT["index_title"])
    _ensure_index_title_style(document)
    index_title.style = "LitMan Analysis Index Title"


def _replace_marker(document: Document, marker: str, value: str):
    paragraph = next((p for p in document.paragraphs if marker in p.text), None)
    if paragraph is None:
        raise ValueError(f"La plantilla no contiene el marcador {marker}")
    for run in paragraph.runs:
        if marker in run.text:
            run.text = run.text.replace(marker, value)
            return paragraph
    raise ValueError(f"El marcador {marker} está fragmentado y no puede sustituirse")


def _populate_toc(document: Document, sections: list[dict], config: dict) -> None:
    marker = "[[INDEX_ENTRIES]]"
    paragraph = next((p for p in document.paragraphs if marker in p.text), None)
    if paragraph is None:
        raise ValueError(f"La plantilla no contiene el marcador {marker}")
    parent = paragraph._element.getparent()
    position = parent.index(paragraph._element)
    parent.remove(paragraph._element)
    for index, section in enumerate(sections, 1):
        entry = _toc_entry(section, index, config, first=index == 1,
                           last=index == len(sections))
        parent.insert(position + index - 1, entry)


def _toc_entry(section: dict, index: int, config: dict, *, first: bool, last: bool):
    paragraph = OxmlElement("w:p")
    properties = OxmlElement("w:pPr")
    widow = OxmlElement("w:widowControl")
    widow.set(qn("w:val"), "0")
    properties.append(widow)
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:leader"), config.get("leader", "none"))
    tab.set(qn("w:pos"), str(config.get("tab_position_twips", 12000)))
    tabs.append(tab)
    properties.append(tabs)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), str(config.get("spacing_before_twips", 60)))
    properties.append(spacing)
    level = int(section.get("level", 1))
    if level > 1:
        indentation = OxmlElement("w:ind")
        indentation.set(qn("w:left"), str((level - 1) * int(config.get("level_indent_twips", 360))))
        properties.append(indentation)
    paragraph.append(properties)
    if first:
        paragraph.append(_field_run("begin", dirty=True))
        instruction = OxmlElement("w:r")
        instruction_text = OxmlElement("w:instrText")
        instruction_text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        instruction_text.text = (' TOC \\h \\u \\z \\t "Heading 1,1,Heading 2,2,'
                                 'Heading 3,3,Heading 4,4,Heading 5,5,Heading 6,6" ')
        instruction.append(instruction_text)
        paragraph.append(instruction)
        paragraph.append(_field_run("separate"))
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), _bookmark_name(section["id"]))
    run = OxmlElement("w:r")
    run.append(_toc_run_properties(config, level))
    text = OxmlElement("w:t")
    text.text = section["title"]
    run.append(text)
    run.append(OxmlElement("w:tab"))
    hyperlink.append(run)
    paragraph.append(hyperlink)
    paragraph.append(_pageref_field(_bookmark_name(section["id"]), config, level))
    if last:
        paragraph.append(_field_run("end"))
    return paragraph


def _toc_run_properties(config: dict, level: int = 1):
    properties = OxmlElement("w:rPr")
    level_config = {**DEFAULT_TOC_LEVELS.get(level, {}),
                    **(config.get("levels") or {}).get(level, {})}
    fonts = OxmlElement("w:rFonts")
    font = level_config.get("font", config.get("font", "Roboto Light"))
    for attribute in ("ascii", "cs", "eastAsia", "hAnsi"):
        fonts.set(qn(f"w:{attribute}"), font)
    properties.append(fonts)
    bold = level_config.get("bold", config.get("bold", True))
    bold_element = OxmlElement("w:b")
    if not bold:
        bold_element.set(qn("w:val"), "0")
    properties.append(bold_element)
    italic = level_config.get("italic", config.get("italic", False))
    italic_element = OxmlElement("w:i")
    if not italic:
        italic_element.set(qn("w:val"), "0")
    properties.append(italic_element)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "none")
    properties.append(underline)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), level_config.get("color", config.get("entry_color", "000000")))
    properties.append(color)
    size = OxmlElement("w:sz")
    font_size = level_config.get("font_size_pt", config.get("font_size_pt", 13))
    size.set(qn("w:val"), str(int(float(font_size) * 2)))
    properties.append(size)
    size_complex = OxmlElement("w:szCs")
    size_complex.set(qn("w:val"), size.get(qn("w:val")))
    properties.append(size_complex)
    return properties


def _pageref_field(bookmark: str, config: dict, level: int = 1):
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), f" PAGEREF {bookmark} \\h ")
    field.set(qn("w:dirty"), "true")
    run = OxmlElement("w:r")
    run.append(_toc_run_properties(config, level))
    field.append(run)
    return field


def _field_run(field_type: str, *, dirty: bool = False):
    run = OxmlElement("w:r")
    field = OxmlElement("w:fldChar")
    field.set(qn("w:fldCharType"), field_type)
    if dirty:
        field.set(qn("w:dirty"), "true")
    run.append(field)
    return run


def _bookmark_name(section_id: str) -> str:
    normalized = re.sub(r"[^A-Za-z0-9_]", "_", section_id)
    return ("litman_core_" + normalized)[:40]


def _add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(1000 + bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(1000 + bookmark_id))
    paragraph._p.insert(1, start)
    paragraph._p.append(end)


def _apply_heading_format(document: Document, paragraph, level: int, config: dict,
                          formal_constraints: dict, numbering_id: int = 2) -> None:
    """Aplica numeración y formato efectivo para importadores como Google Docs."""
    numbering_id = int(config.get("numbering_id", numbering_id))
    properties = paragraph._p.get_or_add_pPr()
    current_numbering = properties.find(qn("w:numPr"))
    if current_numbering is not None:
        properties.remove(current_numbering)
    number_properties = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level - 1))
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), str(numbering_id))
    number_properties.append(ilvl)
    number_properties.append(num_id)
    properties.append(number_properties)

    indentation = _numbering_indentation(document, numbering_id, level - 1)
    if indentation is not None:
        current_indent = properties.find(qn("w:ind"))
        if current_indent is not None:
            properties.remove(current_indent)
        properties.append(indentation)

    style = document.styles[f"Heading {level}"]
    level_config = {**DEFAULT_HEADING_LEVELS.get(level, {}),
                    **(config.get("levels") or {}).get(level, {})}
    font_name = formal_constraints.get("fuente") or level_config.get("font") or style.font.name
    font_size = level_config.get("font_size_pt")
    color = level_config.get("color")
    bold = level_config.get("bold", style.font.bold)
    for run in paragraph.runs:
        if font_name:
            run.font.name = font_name
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_name)
        if font_size is not None:
            run.font.size = Pt(float(font_size))
        elif style.font.size:
            run.font.size = style.font.size
        if bold is not None:
            run.bold = bool(bold)
        if color:
            run.font.color.rgb = _rgb(color)
        elif style.font.color.rgb:
            run.font.color.rgb = style.font.color.rgb


def _numbering_indentation(document: Document, numbering_id: int, level: int):
    numbering = document.part.numbering_part.element
    abstract_id = None
    for number in numbering.findall(qn("w:num")):
        if number.get(qn("w:numId")) == str(numbering_id):
            reference = number.find(qn("w:abstractNumId"))
            abstract_id = reference.get(qn("w:val")) if reference is not None else None
            break
    if abstract_id is None:
        return None
    for abstract in numbering.findall(qn("w:abstractNum")):
        if abstract.get(qn("w:abstractNumId")) != abstract_id:
            continue
        for item in abstract.findall(qn("w:lvl")):
            if item.get(qn("w:ilvl")) != str(level):
                continue
            paragraph_properties = item.find(qn("w:pPr"))
            indentation = paragraph_properties.find(qn("w:ind")) if paragraph_properties is not None else None
            return deepcopy(indentation) if indentation is not None else None
    return None


def _ensure_index_title_style(document: Document) -> None:
    if "LitMan Analysis Index Title" in [style.name for style in document.styles]:
        return
    style = document.styles.add_style("LitMan Analysis Index Title", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = document.styles["Heading 1"]


def _apply_format(document: Document, formatting: dict) -> None:
    font_name = formatting.get("fuente")
    font_size = formatting.get("tamano_pt")
    line_spacing = _line_spacing(formatting.get("interlineado"))
    for style in document.styles:
        if style.type not in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER):
            continue
        if font_name:
            style.font.name = font_name
            style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_name)
        if font_size and style.type == WD_STYLE_TYPE.PARAGRAPH and not style.name.startswith("Heading"):
            style.font.size = Pt(float(font_size))
        if line_spacing and style.type == WD_STYLE_TYPE.PARAGRAPH:
            style.paragraph_format.line_spacing = line_spacing
    _ensure_pending_style(document)
    margins = {
        "top_margin": formatting.get("margen_superior_cm"),
        "bottom_margin": formatting.get("margen_inferior_cm"),
        "left_margin": formatting.get("margen_izquierdo_cm"),
        "right_margin": formatting.get("margen_derecho_cm"),
    }
    for section in document.sections:
        for attribute, value in margins.items():
            if value is not None:
                setattr(section, attribute, Cm(float(value)))


def _ensure_pending_style(document: Document) -> None:
    if "LitMan Analysis Pending" in [style.name for style in document.styles]:
        return
    style = document.styles.add_style("LitMan Analysis Pending", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = next(item for item in document.styles if item.name.lower() == "normal")
    style.font.italic = True


def _ensure_guidance_style(document: Document) -> None:
    if "LitMan Analysis Guidance" in [style.name for style in document.styles]: return
    style = document.styles.add_style("LitMan Analysis Guidance", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = next(item for item in document.styles if item.name.lower() == "normal")
    style.font.size = Pt(9)
    style.font.italic = True


def _insert_editorial_guidance(document, marker, guidance: dict, language: str) -> None:
    if not guidance: return
    _ensure_guidance_style(document)
    labels = UI_TEXT["labels"]
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "EEF5F8"); cell._tc.get_or_add_tcPr().append(shade)
    cell.paragraphs[0].style = "LitMan Analysis Guidance"
    cell.paragraphs[0].add_run(UI_TEXT["guidance"]).bold = True
    pairs = ((labels[0], "purpose"), (labels[1], "evaluation_intent"), (labels[2], "questions_to_answer"), (labels[3], "mandatory_content"), (labels[4], "criteria"), (labels[5], "available_information"), (labels[6], "company_inputs_required"), (labels[7], "evidence_summary"), (labels[8], "avoid"))
    for label, key in pairs:
        value = guidance.get(key)
        if value:
            rendered = _guidance_value(value, language)
            paragraph = cell.add_paragraph(style="LitMan Analysis Guidance"); paragraph.add_run(f"{label}: ").bold = True; paragraph.add_run(rendered)
    if guidance.get("score_confirmed"):
        cell.add_paragraph(f"{labels[9]}: {_format_score(guidance.get('score'), language)}", "LitMan Analysis Guidance")
    if guidance.get("target_pages") is not None:
        cell.add_paragraph(f"{labels[10]}: {guidance['target_pages']} páginas", "LitMan Analysis Guidance")
    elif guidance.get("target_words"):
        cell.add_paragraph(f"{labels[11]}: {guidance['target_words']} palabras", "LitMan Analysis Guidance")
    marker._element.addprevious(table._element)


def _guidance_value(value, language: str = "es") -> str:
    if not isinstance(value, list): return str(value)
    page = "página"
    return "; ".join(f"{x.get('document_id')}, {page} {x.get('page')}" if isinstance(x, dict) else str(x) for x in value)


def _format_score(value, language: str = "es"):
    if value is None: return "pendiente de confirmar"
    number = float(value); rendered = str(int(number)) if number.is_integer() else str(number).replace('.', ',')
    unit = "punto" if number == 1 else "puntos"
    return f"{rendered} {unit}"


def _display_title(section: dict, language: str = "es") -> str:
    return section["title"] if section.get("score") is None else f"{section['title']} ({_format_score(section['score'], language)})"


def _insert_block(document: Document, marker, block: dict, include_internal_pending: bool,
                  table_format: dict) -> None:
    block_type = block["type"]
    if block_type == "paragraph" and block.get("text"):
        _insert_paragraph_before(marker, block["text"])
    elif block_type in ("bullet_list", "numbered_list"):
        style = "List Bullet" if block_type == "bullet_list" else "List Number"
        for item in block.get("items", []):
            _insert_paragraph_before(marker, item, style)
    elif block_type == "pending" and block.get("text") and include_internal_pending:
        _insert_paragraph_before(marker, f"[PENDIENTE] {block['text']}", "LitMan Analysis Pending")
    elif block_type == "table":
        rows = block.get("rows", [])
        headers = block.get("headers", [])
        columns = max(len(headers), max((len(row) for row in rows), default=0))
        if columns:
            table = document.add_table(rows=1 if headers else 0, cols=columns)
            if headers:
                for index, value in enumerate(headers):
                    table.rows[0].cells[index].text = value
            for values in rows:
                cells = table.add_row().cells
                for index, value in enumerate(values):
                    cells[index].text = value
            _format_table(table, bool(headers), table_format)
            marker._element.addprevious(table._element)


def _format_table(table, has_header: bool, config: dict) -> None:
    """Aplica formato OOXML directo; no depende de estilos con nombre."""
    table.autofit = True
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    color = config.get("border_color", "B7B7B7")
    size = str(config.get("border_size", 6))
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)

    margin = str(config.get("cell_margin_twips", 80))
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell_properties = cell._tc.get_or_add_tcPr()
            margins = cell_properties.find(qn("w:tcMar"))
            if margins is None:
                margins = OxmlElement("w:tcMar")
                cell_properties.append(margins)
            for side in ("top", "left", "bottom", "right"):
                side_element = margins.find(qn(f"w:{side}"))
                if side_element is None:
                    side_element = OxmlElement(f"w:{side}")
                    margins.append(side_element)
                side_element.set(qn("w:w"), margin)
                side_element.set(qn("w:type"), "dxa")
            if has_header and row_index == 0:
                shading = cell_properties.find(qn("w:shd"))
                if shading is None:
                    shading = OxmlElement("w:shd")
                    cell_properties.append(shading)
                shading.set(qn("w:val"), "clear")
                shading.set(qn("w:fill"), config.get("header_fill", "71C5E8"))
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.color.rgb = _rgb(config.get("header_text", "000000"))


def _rgb(value: str):
    from docx.shared import RGBColor
    normalized = value.lstrip("#")
    return RGBColor.from_string(normalized if len(normalized) == 6 else "000000")


def _insert_paragraph_before(marker, text: str, style: str | None = None):
    paragraph = marker.insert_paragraph_before(text)
    if style:
        try:
            paragraph.style = style
        except KeyError:
            pass
    return paragraph


def _mark_fields_for_update(document: Document) -> None:
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def _line_spacing(value) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    normalized = str(value).lower().replace(",", ".").strip()
    aliases = {"sencillo": 1.0, "simple": 1.0, "doble": 2.0}
    if normalized in aliases:
        return aliases[normalized]
    try:
        return float(normalized)
    except ValueError:
        return None
