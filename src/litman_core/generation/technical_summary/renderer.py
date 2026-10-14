from __future__ import annotations
from pathlib import Path
from typing import Any
from docx import Document
from ..docx_helpers import MISSING, abbreviate_text, add_callout, add_data_table, add_empty_message, add_fact_table, amount_value, configure_document, display_value, named_description


def _items(value: Any) -> list[Any]: return value if isinstance(value, list) else []
def _compact(value: Any, limit: int = 180) -> str: return abbreviate_text(value, limit)
def _header(value: dict[str, Any]) -> str:
    return " · ".join(str(x) for x in (value.get("expediente"), value.get("nombre")) if x not in (None, "")) or MISSING


def write_technical_summary(data: dict[str, Any], path: Path) -> None:
    doc = Document(); header = data.get("cabecera") if isinstance(data.get("cabecera"), dict) else {}
    configure_document(doc, title="Resumen técnico", subject="Ficha técnica del expediente", document_label="Resumen técnico", header_text=_header(header))
    doc.add_heading("RESUMEN TÉCNICO", 1)
    doc.add_paragraph(display_value(header.get("nombre")), style="LitMan Lead")
    identity = " · ".join(x for x in (display_value(header.get("organismo")), f"Expediente {header['expediente']}" if header.get("expediente") else "") if x)
    doc.add_paragraph(identity or MISSING, style="LitMan Metadata")
    amount = amount_value(header.get("presupuesto_sin_iva"), header.get("moneda"))
    phases = len(_items(data.get("fases"))) or MISSING
    add_fact_table(doc, ["EXPEDIENTE", "FECHA LÍMITE", "PRESUPUESTO", "FASES"], [header.get("expediente"), header.get("fecha_limite"), amount, phases])
    _context(doc, data); _objectives(doc, data.get("objetivos")); _phases(doc, data.get("fases")); _resources(doc, data); _key_aspects(doc, data)
    path.parent.mkdir(parents=True, exist_ok=True); doc.save(path)


def _context(doc: Document, data: dict[str, Any]) -> None:
    doc.add_heading("1. Contexto", 2)
    doc.add_paragraph(_compact(data.get("contexto")) if data.get("contexto") else MISSING, style="LitMan Callout")
    scope = "; ".join(display_value(x) for x in _items(data.get("alcance"))) or MISSING
    work = "; ".join(display_value(x) for x in _items(data.get("trabajos"))) or MISSING
    add_data_table(doc, ["ELEMENTO", "RESUMEN"], [("Alcance", _compact(scope)), ("Trabajos principales", _compact(work))], [4.2, 12.8])


def _objectives(doc: Document, value: Any) -> None:
    doc.add_heading("2. Objetivos", 2); values = _items(value)
    if not values: add_empty_message(doc, "No consta información sobre objetivos."); return
    for index, item in enumerate(values, 1): doc.add_paragraph(f"{index}. {_compact(item, 160)}", style="LitMan Compact")


def _phases(doc: Document, value: Any) -> None:
    doc.add_heading("3. Fases y entregables", 2); values = _items(value)
    if not values: add_empty_message(doc, "No consta información sobre fases."); return
    rows = []
    for raw in values:
        item = raw if isinstance(raw, dict) else {}
        deliverables = _items(item.get("entregables"))
        if not deliverables: deliverables = [MISSING]
        for deliverable in deliverables:
            if isinstance(deliverable, dict):
                name = display_value(deliverable.get("nombre"))
                description = deliverable.get("descripcion")
                deliverable_text = name if not description else f"{name} — {_compact(description)}"
            else:
                deliverable_text = display_value(deliverable)
            rows.append((display_value(item.get("nombre")), _compact(item.get("descripcion")), display_value(item.get("plazo")), deliverable_text))
    add_data_table(doc, ["FASE", "ACTIVIDAD / DESCRIPCIÓN", "PLAZO", "ENTREGABLES"], rows, [3.0, 6.0, 2.5, 5.5])


def _resources(doc: Document, data: dict[str, Any]) -> None:
    doc.add_heading("4. Tecnologías y recursos", 2)
    technologies = "; ".join(display_value(x) for x in _items(data.get("tecnologias"))) or MISSING
    profiles = "; ".join(named_description(x) for x in _items(data.get("equipo"))) or MISSING
    add_data_table(doc, ["TECNOLOGÍAS Y PLATAFORMAS", "PERFILES / RECURSOS"], [(_compact(technologies, 240), _compact(profiles, 240))], [8.0, 9.0])


def _key_aspects(doc: Document, data: dict[str, Any]) -> None:
    doc.add_heading("5. Aspectos clave", 2)
    requirements = [x for x in _items(data.get("requisitos_oferta")) if isinstance(x, dict)]
    if requirements:
        doc.add_paragraph("Requisitos", style="LitMan Card Title")
        for item in requirements:
            label = "OBLIGATORIO" if item.get("obligatorio") is True else "NO OBLIGATORIO" if item.get("obligatorio") is False else "OBLIGATORIEDAD NO INDICADA"
            add_callout(doc, label + " · " + display_value(item.get("tipo")), _compact(item.get("descripcion")))
    else: add_empty_message(doc, "No constan requisitos de la oferta.")
    criteria = [x for x in _items(data.get("criterios_adjudicacion")) if isinstance(x, dict)]
    if criteria:
        doc.add_paragraph("Criterios de adjudicación", style="LitMan Card Title")
        for item in criteria:
            points = f"{item['puntuacion_maxima']} puntos" if item.get("puntuacion_maxima") is not None else "Puntuación no indicada"
            add_callout(doc, display_value(item.get("nombre")) + " · " + points, _compact(item.get("descripcion")))
    else: add_empty_message(doc, "No constan criterios de adjudicación.")
    quality = data.get("calidad") if isinstance(data.get("calidad"), dict) else {}; limits = _items(quality.get("limitaciones"))
    if limits: add_callout(doc, "ALERTA: LIMITACIONES DEL ANÁLISIS", "; ".join(display_value(x) for x in limits), kind="warning")
