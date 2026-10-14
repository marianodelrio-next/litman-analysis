from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from litman_core.application.ports.llm import LLMResult
from litman_core.config import Settings
from litman_core.pipelines.assessment.service import AssessmentPipeline
from litman_core.pipelines.detailed_analysis.service import DetailedAnalysisPipeline

from conftest import dossier


QUOTE = "The methodology must be described."
PRICE_FORMULA = "Score = lowest bid / evaluated bid * 10"
ABNORMAL_FORMULA = "R=(85%+C×5%)×average bid"


class ScriptedSpanishProvider:
    """Test double: its responses are authored Spanish fixtures, never a translation model."""

    def __init__(self) -> None:
        self.instructions: list[str] = []

    def generate(self, *, instructions, prompt, schema_path, case_dir, output_paths):
        self.instructions.append(instructions)
        document_id = re.search(r'"documento_id":\s*"([^"]+)"', prompt).group(1)
        data = _assessment(document_id) if "primera valoración" in instructions else _detailed(document_id)
        return LLMResult(data, "scripted", "fixture", "fixture-response", {}, {})


def _assessment(document_id: str) -> dict:
    return {
        "control_documental": {"documentos": [{"documento_id": document_id, "tipo": "pliego técnico", "fase": "licitacion", "procesado": True, "observaciones": "Documento revisado."}], "ausencias_relevantes": []},
        "ficha": {"objeto_contrato": "Servicio técnico", "numero_expediente": None, "organismo": None,
                  "presupuesto": {"base_sin_iva": None, "iva": None, "base_con_iva": None, "valor_estimado": None, "moneda": None},
                  "fecha_limite_presentacion": None, "duracion_contrato": None, "certificaciones_exigidas": [],
                  "equipo": {"numero_personas_exigido": None, "perfiles": [], "otros_colectivos": []},
                  "stack_tecnologico": {"tecnologias": [], "nuevo_desarrollo": None, "stack_nuevo_desarrollo": []},
                  "presencialidad": {"exigida": None, "alcance": None, "frecuencia": None, "modalidad": "no_consta"},
                  "provincia": None, "subcontratacion": None, "solvencia_tecnica": None, "solvencia_economica": None,
                  "tipo_contrato": None, "tipo_procedimiento": None, "tramitacion_urgente": None, "cpv": [],
                  "lotes": {"existen": None, "descripcion": None}, "resumen": "El pliego exige describir la metodología."},
        "clasificacion": {"tipologias": ["Servicio técnico"], "tecnologias": [], "sectores": [], "senales_positivas": ["Requisito localizado."], "senales_riesgo": []},
        "calidad_informacion": {"estado": "suficiente", "limitaciones": [], "contradicciones": []},
        "evidencias": [{"campo": "ficha.objeto_contrato", "documento_id": document_id, "pagina": 1, "seccion": None, "texto_citado": QUOTE}],
    }


def _detailed(document_id: str) -> dict:
    value = dossier()
    value["contexto"] = "El expediente requiere una metodología verificable."
    value["memoria_tecnica"]["idioma"] = "en"
    value["memoria_tecnica"]["secciones"][0]["titulo"] = "Planteamiento del proyecto"
    value["memoria_tecnica"]["secciones"][0]["descripcion"] = "Descripción técnica exacta en español."
    value["evidencias"] = [{"campo": "metodologia", "documento_id": document_id, "pagina": 1, "texto_citado": QUOTE}]
    value["simulador_economico"]["formula_precio"]["expresion_literal"] = PRICE_FORMULA
    value["simulador_economico"]["baja_anormal"]["formula_literal"] = ABNORMAL_FORMULA
    source_evidence = {"campo": "metodologia", "documento_id": document_id, "pagina": 1, "texto_citado": QUOTE}
    value["simulador_economico"]["formula_precio"]["evidencia"] = source_evidence
    value["simulador_economico"]["criterios_no_precio"][0]["evidencia"] = source_evidence
    return value


def test_mixed_english_catalan_source_runs_real_stages_with_spanish_scripted_outputs(tmp_path: Path) -> None:
    case = tmp_path / "case"; case.mkdir()
    source = Path(__file__).parents[1] / "fixtures/spanish-output-language/english-source.md"
    (case / "tender.md").write_text(source.read_text(encoding="utf-8"), encoding="utf-8")
    assert "Plantejament del projecte" in (case / "tender.md").read_text(encoding="utf-8")
    settings = Settings(runs_root=tmp_path / "runs")
    provider = ScriptedSpanishProvider()
    assessment_run, assessment = AssessmentPipeline(settings, provider).execute(case, case_id="english")
    assert assessment["validation"]["valid"]
    assessment_json = json.loads((assessment_run / "output/extraction.json").read_text())
    assert assessment_json["ficha"]["resumen"].startswith("El pliego exige")
    assert assessment["report"]
    report_text = (assessment_run / "output/report.json").read_text(encoding="utf-8")
    html_text = (assessment_run / "output/ficha_oportunidad.html").read_text(encoding="utf-8")
    for report in (report_text, html_text):
        assert "El pliego exige describir la metodología" in report
        assert QUOTE in report
    detailed_run, result = DetailedAnalysisPipeline(settings, provider).execute(case, case_id="english", assessment_result=assessment)
    detailed = json.loads((detailed_run / "output/dossier.json").read_text())
    evidence = json.loads((assessment_run / "output/evidence-index.json").read_text())
    assert detailed["contexto"].startswith("El expediente requiere")
    assert detailed["memoria_tecnica"]["secciones"][0]["titulo"] == "Planteamiento del proyecto"
    assert evidence["evidences"][0]["verification"] == "verified"
    assert detailed["evidencias"][0]["texto_citado"] == QUOTE
    assert detailed["simulador_economico"]["formula_precio"]["expresion_literal"] == PRICE_FORMULA
    assert detailed["simulador_economico"]["baja_anormal"]["formula_literal"] == ABNORMAL_FORMULA
    plan = json.loads((detailed_run / "input/document_plan.json").read_text())
    assert plan["sections"][0]["title"] == "Planteamiento del proyecto"
    assert plan["sections"][0]["purpose"] == "Descripción técnica exacta en español."
    assert "Apartado técnico" not in json.dumps(plan, ensure_ascii=False)
    assert all("campos semánticos y narrativos" in text or "contenido visible" in text for text in provider.instructions[:2])
    documents = result["generation"]["documents"]
    text = "\n".join("\n".join(p.text for p in Document(path).paragraphs) for path in documents if path.endswith(".docx"))
    assert "Propuesta técnica" in text and "Planteamiento del proyecto" in text and detailed["contexto"] in text and QUOTE not in text
    workbook = load_workbook(next(path for path in documents if path.endswith(".xlsx")), data_only=False)
    cells = "\n".join(str(cell.value) for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row if cell.value is not None)
    assert PRICE_FORMULA in cells
