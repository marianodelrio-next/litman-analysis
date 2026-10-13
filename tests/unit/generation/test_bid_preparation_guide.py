import json
import copy
from pathlib import Path

from docx import Document
from docx.shared import Cm

from conftest import dossier
from litman_core.generation.bid_preparation_guide import BidPreparationGuideGenerator
from litman_core.generation.bid_preparation_guide.renderer import write_bid_preparation_guide


def _text(path: Path) -> str:
    doc = Document(path)
    return "\n".join([p.text for p in doc.paragraphs] + [cell.text for table in doc.tables for row in table.rows for cell in row.cells])


def test_bid_guide_contains_operational_data_nulls_pending_and_evidence(tmp_path: Path) -> None:
    output = tmp_path / "guia.docx"; write_bid_preparation_guide(dossier(), output)
    text = _text(output)
    for value in ("GUÍA DE PREPARACIÓN DE LA OFERTA", "Documentación que presentar", "SOBRE", "Memoria técnica", "Metodología", "No consta", "No incluir precio", "Documento técnico", "Información económica", "Resumen y coordinación", "Generales", "Falta información", "PPT", "Página", "Calidad y limitaciones"):
        assert value in text
    assert "None" not in text and "Entregables del contrato" not in text


def test_bid_guide_uses_suggested_structure_when_imposed_sections_are_absent(tmp_path: Path) -> None:
    data = dossier(memoria_tecnica={"secciones": [], "formato": {}}, estructura_documento_tecnico=["Propuesta sugerida"])
    output = tmp_path / "guide.docx"; write_bid_preparation_guide(data, output)
    text = _text(output)
    assert "ESTRUCTURA SUGERIDA POR EL ANÁLISIS" in text and "Propuesta sugerida" in text


def test_bid_guide_rerenders_snapshot_with_its_stable_contract(tmp_path: Path) -> None:
    run = tmp_path / "run"; (run / "input").mkdir(parents=True); (run / "output").mkdir()
    data = dossier(); (run / "input" / "dossier.json").write_text(json.dumps(data), encoding="utf-8")
    result = BidPreparationGuideGenerator().rerender(run)
    assert Path(result["documents"][0]).name == "guia_preparacion_oferta.docx"
    assert result["id"] == "bid-preparation-guide"
    assert result["generator_version"] == "bid-preparation-guide-generator-v3"
    assert result["pending"] == data["pendientes"] and result["metadata"]["format"] == "docx"


def test_bid_guide_preserves_budget_and_groups_interleaved_evidence(tmp_path: Path) -> None:
    evidence = [
        {"campo": "A", "documento_id": "one", "pagina": 1, "texto_citado": "first"},
        {"campo": "B", "documento_id": "two", "pagina": 2, "texto_citado": "second"},
        {"campo": "A", "documento_id": "three", "pagina": None, "texto_citado": "third"},
    ]
    data = dossier(evidencias=evidence)
    data["cabecera"]["presupuesto_sin_iva"] = 200000
    output = tmp_path / "guide.docx"; write_bid_preparation_guide(data, output)
    doc = Document(output); text = _text(output)
    assert "200000 EUR" in text
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 3"]
    assert headings.count("A") == 1 and headings.count("B") == 1
    assert all(item in text for item in ("one", "two", "three", "Página: No consta"))


def test_bid_guide_marks_unknown_format_without_a_checkbox(tmp_path: Path) -> None:
    data = dossier(memoria_tecnica={"formato": {}, "portada_computa_limite": None, "indice_computa_limite": False, "anexos_computan_limite": True})
    output = tmp_path / "format.docx"; write_bid_preparation_guide(data, output)
    text = _text(output)
    assert "REVISAR · No consta el tamaño de fuente." in text
    assert "□ REVISAR" not in text
    assert "□ El índice no computa en el límite." in text
    assert "□ Los anexos computan en el límite." in text
    assert "REVISAR · No consta el tamaño de fuente.\nNo consta" not in text


def test_bid_guide_has_accessible_document_structure(tmp_path: Path) -> None:
    output = tmp_path / "structure.docx"; write_bid_preparation_guide(dossier(), output)
    doc = Document(output); section = doc.sections[0]
    tolerance = Cm(.01)
    assert abs(section.page_width - Cm(21)) < tolerance and abs(section.page_height - Cm(29.7)) < tolerance
    assert abs(section.left_margin - Cm(2)) < tolerance and abs(section.right_margin - Cm(2)) < tolerance
    assert doc.styles["Normal"].font.name == "Arial"
    assert sum(p.style.name == "Heading 1" for p in doc.paragraphs) == 1
    assert any(p._p.xpath('.//w:br[@w:type="page"]') for p in doc.paragraphs)
    assert section.header.paragraphs[0].text
    assert section.footer._element.xpath('.//w:fldSimple[contains(@w:instr,"PAGE")]')
    assert all(len(table.columns) <= 4 for table in doc.tables)
    for table in doc.tables:
        assert table.rows[0]._tr.xpath('./w:trPr/w:tblHeader')
        assert all(row._tr.xpath('./w:trPr/w:cantSplit') for row in table.rows)


def test_bid_guide_keeps_input_and_all_rule_and_evidence_information(tmp_path: Path) -> None:
    data = dossier()
    data["memoria_tecnica"]["secciones"][0]["puntuacion"] = None
    data["memoria_tecnica"]["secciones"][0]["limite_paginas"] = None
    data["memoria_tecnica"]["reglas"].append({"tipo": "otra", "descripcion": "Regla especial", "consecuencia": "Revisar manualmente", "evidencia": None})
    data["evidencias"].append({"campo": "larga", "documento_id": "DOC", "pagina": None, "texto_citado": "x" * 600})
    original = copy.deepcopy(data)
    output = tmp_path / "complete.docx"; write_bid_preparation_guide(data, output)
    text = _text(output)
    assert data == original
    assert "Puntuación no indicada" in text and "Límite no indicado" in text
    assert "Regla especial" in text and "Revisar manualmente" in text
    assert "[cita abreviada]" in text and "Página: No consta" in text
