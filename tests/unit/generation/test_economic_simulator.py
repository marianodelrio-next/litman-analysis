import json
import zipfile
from copy import deepcopy
from pathlib import Path

from openpyxl import load_workbook

from litman_core.generation.service import generate_documents
from litman_core.generation.economic_simulator.calculator import calculate_scenario
from litman_core.generation.economic_simulator.generator import _pending, render_existing_economic_run
from litman_core.generation.economic_simulator.planner import build_economic_plan, skeleton_economic_inputs
from litman_core.generation.economic_simulator.renderer import render_economic_workbook
from litman_core.resource_paths import RESOURCE_ROOT

from conftest import dossier


RESOURCES = RESOURCE_ROOT


def _multiple_criteria_dossier() -> dict:
    value = deepcopy(dossier())
    value["simulador_economico"]["criterios_no_precio"] = [
        {"nombre": "Metodología", "tipo": "juicio_valor", "puntuacion_maxima": 50,
         "valor_minimo": 0, "valor_maximo": 50, "unidad": "puntos",
         "descripcion": "Calidad de la metodología", "evidencia": None},
        {"nombre": "Equipo", "tipo": "automatico_cuantitativo", "puntuacion_maxima": 40,
         "valor_minimo": 0, "valor_maximo": 40, "unidad": "puntos",
         "descripcion": "Experiencia y composición", "evidencia": None},
    ]
    return value


def test_economic_simulator_uses_one_technical_column_and_is_rerenderable(tmp_path: Path) -> None:
    dossier_path = tmp_path / "case" / "detailed-analysis" / "run-1" / "output" / "dossier.json"
    dossier_path.parent.mkdir(parents=True)
    dossier_path.write_text(json.dumps(_multiple_criteria_dossier(), ensure_ascii=False), encoding="utf-8")
    run_dir = tmp_path / "generation-run"
    manifest = generate_documents(dossier_path, output_dir=run_dir, artifacts=["economic-simulator"], provider=None)
    output = run_dir / "output" / "simulador_economico.xlsx"

    assert output.exists()
    assert manifest["economic_generation"]["template_version"] == "1.1.0"
    workbook = load_workbook(output, data_only=False)
    assert workbook.sheetnames == ["Simulador", "Criterios y ayuda", "Comprobaciones", "Pendientes"]
    simulator = workbook["Simulador"]
    headers = [simulator.cell(14, column).value for column in range(1, 11)]
    assert headers == ["Licitador", "Precio ofertado", "% bajada", "Puntuación técnica\n(0–90 p)",
                       "Puntos precio", "Puntuación total", "C cualitativo",
                       "Umbral baja anormal", "Estado", "Posición"]
    assert "Metodología" not in "\n".join(str(value) for value in headers)
    assert "Equipo" not in "\n".join(str(value) for value in headers)
    assert simulator.max_column == 10
    assert simulator["D15"].fill.fgColor.rgb.endswith("FFF2CC")
    assert simulator["F15"].value == '=IF(OR(B15="",D15=""),"",SUM(E15,D15))'
    assert simulator["G15"].value == '=IF(OR(B15="",D15=""),"",IFERROR(D15/$G$7,0))'
    validations = [item for item in simulator.data_validations.dataValidation if "D15" in str(item.sqref)]
    assert len(validations) == 1
    assert validations[0].type == "decimal"
    assert validations[0].formula1 == "0" and validations[0].formula2 == "90"
    criteria = workbook["Criterios y ayuda"]
    assert [criteria.cell(row, 1).value for row in (6, 7)] == ["Metodología", "Equipo"]

    inputs_path = run_dir / "output" / "economic_inputs.json"
    inputs = json.loads(inputs_path.read_text(encoding="utf-8"))
    assert inputs["version"] == "economic-inputs-v2"
    assert inputs["scenarios"][0]["bidders"][0] == {"name": "NEXT", "price": None, "technical_score": None}
    pending = json.loads((run_dir / "output" / "economic_pending.json").read_text(encoding="utf-8"))
    assert any(item["description"] == "Definir la puntuación técnica esperada de NEXT." for item in pending)
    inputs["scenarios"][0]["bidders"][0].update(price=150000, technical_score=90)
    inputs_path.write_text(json.dumps(inputs, ensure_ascii=False), encoding="utf-8")
    result = render_existing_economic_run(run_dir)
    assert result["validation"]["valid"]
    assert result["validation"]["calculations"]["scenario-01"][0]["total_points"] == 100


def test_aggregated_score_calculates_total_threshold_and_zero() -> None:
    plan = build_economic_plan(_multiple_criteria_dossier())
    scenario = {"bidders": [
        {"name": "NEXT", "price": 150000, "technical_score": 90},
        {"name": "A", "price": 180000, "technical_score": 0},
        {"name": "B", "price": 190000, "technical_score": 50},
        {"name": "C", "price": 200000, "technical_score": 40},
    ]}
    result = calculate_scenario(plan, scenario)
    assert plan["technical_max_points"] == 90
    assert result[0]["price_points"] == 10
    assert result[0]["qualitative_points"] == 90
    assert result[0]["total_points"] == 100
    assert result[0]["possible_abnormal"] is True
    assert result[1]["qualitative_points"] == 0
    assert result[1]["total_points"] is not None


def test_validation_rejects_invalid_aggregate_scores_and_accepts_null(tmp_path: Path) -> None:
    plan = build_economic_plan(_multiple_criteria_dossier())
    template = RESOURCES / "templates" / "economic-simulator" / "standard-v1" / "template.xlsx"
    for score, expected in [(-1, "fuera de rango"), (91, "fuera de rango"), (True, "debe ser numérica"), (None, None)]:
        inputs = skeleton_economic_inputs(plan, ["Prueba"], 0)
        inputs["scenarios"][0]["bidders"][0]["technical_score"] = score
        validation = render_economic_workbook(template, plan, inputs, tmp_path / f"{score}.xlsx")
        if expected:
            assert not validation["valid"]
            assert any(expected in error for error in validation["errors"])
        else:
            assert validation["valid"]


def test_no_technical_criteria_generates_valid_zero_only_input(tmp_path: Path) -> None:
    value = _multiple_criteria_dossier()
    value["simulador_economico"]["criterios_no_precio"] = []
    plan = build_economic_plan(value)
    inputs = skeleton_economic_inputs(plan, ["Prueba"], 0)
    assert inputs["scenarios"][0]["bidders"][0]["technical_score"] == 0
    assert not any("puntuación técnica" in item["description"] for item in _pending(plan, inputs))
    template = RESOURCES / "templates" / "economic-simulator" / "standard-v1" / "template.xlsx"
    output = tmp_path / "no-criteria.xlsx"
    assert render_economic_workbook(template, plan, inputs, output)["valid"]
    workbook = load_workbook(output)
    technical_validation = next(item for item in workbook["Simulador"].data_validations.dataValidation
                                if "D15" in str(item.sqref))
    assert technical_validation.type == "list"
    assert technical_validation.formula1 == '"0"'


def test_economic_plan_normalizes_percentage_bases() -> None:
    value = dossier()
    value["simulador_economico"]["baja_anormal"]["reglas"][0]["base"] = 85
    value["simulador_economico"]["baja_anormal"]["maximo_cualitativo"] = 93
    plan = build_economic_plan(value)
    assert plan["abnormally_low"]["reglas"][0]["base"] == 0.85
    assert any("normalizada" in warning for warning in plan["warnings"])
    assert any("suman 90" in warning and "usa 93" in warning for warning in plan["warnings"])


def test_rerender_accepts_legacy_v1_criterion_scores(tmp_path: Path) -> None:
    dossier_path = tmp_path / "case" / "detailed-analysis" / "run-1" / "output" / "dossier.json"
    dossier_path.parent.mkdir(parents=True)
    dossier_path.write_text(json.dumps(_multiple_criteria_dossier()), encoding="utf-8")
    run_dir = tmp_path / "generation-run"
    generate_documents(dossier_path, output_dir=run_dir, artifacts=["economic-simulator"], provider=None)
    inputs_path = run_dir / "output" / "economic_inputs.json"
    inputs = json.loads(inputs_path.read_text(encoding="utf-8"))
    inputs["version"] = "economic-inputs-v1"
    for scenario in inputs["scenarios"]:
        for bidder in scenario["bidders"]:
            bidder.pop("technical_score")
            bidder["criterion_scores"] = {"criterion-001": 50, "criterion-002": 40}
    inputs["scenarios"][0]["bidders"][0]["price"] = 150000
    inputs_path.write_text(json.dumps(inputs), encoding="utf-8")
    result = render_existing_economic_run(run_dir)
    assert result["validation"]["valid"]
    workbook = load_workbook(run_dir / "output" / "simulador_economico.xlsx", data_only=False)
    assert workbook["Simulador"]["D15"].value == 90


def test_economic_template_is_neutral() -> None:
    template = RESOURCES / "templates" / "economic-simulator" / "standard-v1" / "template.xlsx"
    with zipfile.ZipFile(template) as archive:
        xml = "\n".join(archive.read(name).decode("utf-8", errors="ignore")
                          for name in archive.namelist() if name.endswith(".xml"))
    for term in ("UIB", "21/26", "NEXT", "DAMA", "206611"):
        assert term not in xml
