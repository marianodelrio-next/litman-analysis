import json
import zipfile
from xml.etree import ElementTree
from pathlib import Path

from docx import Document
from jsonschema import Draft202012Validator

from litman_analysis.generation.technical_proposal.content import fallback_content, skeleton_content
from litman_analysis.generation.technical_proposal.planner import build_document_plan, normalize_output_language
from litman_analysis.generation.technical_proposal.renderer import render_technical_document
from litman_analysis.generation.service import generate_documents, render_existing_run
from litman_analysis.resource_paths import RESOURCE_ROOT

from conftest import dossier
RESOURCES = RESOURCE_ROOT


def test_detailed_analysis_fixture_matches_schema() -> None:
    schema = json.loads((RESOURCES / "schemas" / "detailed-analysis.schema.json").read_text())
    assert list(Draft202012Validator(schema).iter_errors(dossier())) == []
    assert schema["properties"]["schema_version"] == {
        "type": "string", "const": "detailed-analysis-v1",
    }


def test_plan_preserves_tender_order_and_constraints() -> None:
    plan = build_document_plan(dossier())
    assert plan["structure_source"] == "tender"
    assert [item["title"] for item in plan["sections"]] == ["Metodología"]
    assert plan["sections"][0]["requirements"] == ["Explicar la metodología"]
    assert plan["formal_constraints"]["fuente"] == "Arial"
    assert plan["word_budget"]["total_target_words"] <= 3500
    assert plan["sections"][0]["evidence_fields"] == []


def test_generic_criterion_destination_does_not_contaminate_other_sections() -> None:
    value = dossier()
    value["criterios_adjudicacion"].append({
        "nombre": "Equipo", "puntuacion_maxima": 5, "tipo": "juicio",
        "donde_responder": "Memoria técnica", "descripcion": None,
    })
    plan = build_document_plan(value)
    assert plan["sections"][0]["criteria"] == ["Metodología"]


def test_language_normalization_prefers_spanish_for_multilingual_option() -> None:
    assert normalize_output_language("Catalán o español") == "es"
    assert normalize_output_language("English") == "en"
    assert normalize_output_language(None) == "es"


def test_old_dossier_gets_safe_fallback_plan() -> None:
    old = dossier()
    old.pop("memoria_tecnica")
    plan = build_document_plan(old)
    assert plan["structure_source"] == "fallback"
    assert plan["pending_decisions"]


def test_plan_uses_distinct_contextual_subsections_and_fallback_score() -> None:
    value = dossier()
    value["memoria_tecnica"]["secciones"] = []
    value["estructura_documento_tecnico"] = ["Metodología", "Equipo requerido"]
    value["criterios_adjudicacion"] = [
        {"nombre": "Metodología", "puntuacion_maxima": 5, "tipo": "juicio", "donde_responder": "Metodología", "descripcion": None},
        {"nombre": "Equipo requerido", "puntuacion_maxima": 1, "tipo": "juicio", "donde_responder": "Equipo requerido", "descripcion": None},
    ]
    plan = build_document_plan(value)
    assert [x["title"] for x in plan["sections"][0]["subsections"]] == ["Principios metodológicos", "Organización de las actividades", "Validación y control de calidad", "Seguimiento y mejora continua"]
    assert [x["title"] for x in plan["sections"][1]["subsections"]] == ["Organización del equipo", "Perfiles y responsabilidades", "Modelo de coordinación", "Disponibilidad y continuidad pendiente de confirmar"]
    assert plan["sections"][0]["display_title"] == "Metodología (5 puntos)"
    assert plan["sections"][1]["display_title"] == "Equipo requerido (1 punto)"
    assert plan["sections"][0]["score_source"] == "criterion"


def test_short_filename_preserves_acronyms_and_meaningful_terms(tmp_path: Path) -> None:
    value = dossier()
    value["cabecera"].update(nombre="Aplicaciones de Inteligencia Artificial (IA) aplicada a expedientes de Compras (CARMEN)", expediente="5100073845")
    dossier_path = tmp_path / "dossier.json"; dossier_path.write_text(json.dumps(value), encoding="utf-8")
    manifest = generate_documents(dossier_path, output_dir=tmp_path / "run", artifacts=["technical-document"], provider=None)
    assert Path(manifest["documents"][0]).name == "Propuesta técnica 5100073845 Inteligencia Artificial IA aplicada expedientes Compras CARMEN.docx"


def test_english_plan_localizes_generated_structure_and_nested_evidence() -> None:
    value = dossier()
    value["memoria_tecnica"]["idioma"] = "en"
    section = value["memoria_tecnica"]["secciones"][0]
    section.update(titulo="Methodology", descripcion=None, evidencias=[{
        "campo": "methodology", "documento_id": "PPT", "pagina": 12,
        "texto_citado": "The methodology must be described.",
    }])
    plan = build_document_plan(value)
    planned = plan["sections"][0]
    assert planned["display_title"] == "Methodology (20 points)"
    assert planned["subsections"][0]["title"] == "Methodological principles"
    assert planned["guidance"]["evidence_summary"][0]["document_id"] == "PPT"
    assert "Explain verifiably" in planned["guidance"]["purpose"]


def test_filename_does_not_repeat_compound_reference(tmp_path: Path) -> None:
    value = dossier()
    value["cabecera"].update(nombre="Servicio de análisis EXP-1", expediente="EXP-1")
    dossier_path = tmp_path / "dossier.json"; dossier_path.write_text(json.dumps(value), encoding="utf-8")
    manifest = generate_documents(dossier_path, output_dir=tmp_path / "run", artifacts=["technical-document"], provider=None)
    assert Path(manifest["documents"][0]).name == "Propuesta técnica EXP-1 análisis.docx"


def test_identity_pending_and_guidance_toggle_are_rendered_from_plan(tmp_path: Path) -> None:
    value = dossier(); value["cabecera"].update(nombre=None, expediente=None)
    dossier_path = tmp_path / "dossier.json"; dossier_path.write_text(json.dumps(value), encoding="utf-8")
    run = tmp_path / "run"
    generate_documents(dossier_path, output_dir=run, artifacts=["technical-document"], provider=None)
    pending = json.loads((run / "output/pending.json").read_text(encoding="utf-8"))
    assert sum("nombre ni expediente" in item["description"] for item in pending) == 1
    result = render_existing_run(run, include_editorial_guidance=False)
    clean = Document(result["document"])
    assert not any("Guía de redacción" in p.text for p in clean.paragraphs)
    assert all("EEF5F8" not in str(table._tbl.xml) for table in clean.tables)


def test_render_uses_clean_template_and_keeps_dynamic_section(tmp_path: Path) -> None:
    plan = build_document_plan(dossier())
    content = fallback_content(dossier(), plan)
    content["sections"][0]["blocks"].append({
        "type": "table", "text": None, "items": [],
        "headers": ["Fase", "Plazo"], "rows": [["Diseño", "2 semanas"]],
    })
    output = tmp_path / "proposal.docx"
    validation = render_technical_document(
        RESOURCES / "templates" / "technical-proposal" / "next-standard-v1" / "template.docx",
        content, plan, output,
    )
    assert validation["valid"]
    text = "\n".join(item.text for item in Document(output).paragraphs)
    assert "1. Metodología" in text
    assert "OCS" not in text
    assert "[[DOCUMENT_BODY]]" not in text
    assert "Propuesta técnica" in text
    assert "Índice" in text
    assert "Technical proposal" not in text
    assert "1. Metodología" in text
    assert validation["page_limit_validation"] == "not_performed"
    rendered = Document(output)
    assert [paragraph.text for paragraph in rendered.paragraphs[:8]] == [
        "", "", "", "", "", "", "Propuesta técnica", "Servicio de analítica\nExpte. EXP-1",
    ]
    assert rendered.tables
    shading = rendered.tables[0].rows[0].cells[0]._tc.tcPr.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd"
    )
    assert shading is not None
    assert shading.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill") == "EEF5F8"

    source = RESOURCES / "templates" / "technical-proposal" / "next-standard-v1" / "template.docx"
    with zipfile.ZipFile(source) as template_archive, zipfile.ZipFile(output) as output_archive:
        assert output_archive.read("word/header1.xml") == template_archive.read("word/header1.xml")
        assert output_archive.read("word/media/image25.png") == template_archive.read("word/media/image25.png")
        root = ElementTree.fromstring(output_archive.read("word/document.xml"))
    word = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    toc = [node.text or "" for node in root.iter(word + "instrText") if "TOC" in (node.text or "")]
    pagerefs = [node for node in root.iter(word + "fldSimple") if "PAGEREF" in node.get(word + "instr", "")]
    bookmarks = [node.get(word + "name") for node in root.iter(word + "bookmarkStart")
                 if (node.get(word + "name") or "").startswith("litman_analysis_")]
    blue_runs = [node for node in root.iter(word + "color") if node.get(word + "val") == "4FA6DC"]
    assert toc
    assert len(pagerefs) >= 1
    assert len(bookmarks) >= 1
    assert blue_runs


def test_render_localizes_template_interface_in_english(tmp_path: Path) -> None:
    plan = build_document_plan(dossier())
    plan["language"] = "en"
    content = fallback_content(dossier(), plan)
    output = tmp_path / "proposal-en.docx"
    render_technical_document(
        RESOURCES / "templates" / "technical-proposal" / "next-standard-v1" / "template.docx",
        content, plan, output,
    )
    text = "\n".join(item.text for item in Document(output).paragraphs)
    assert "Technical proposal" in text
    assert "Contents" in text
    assert "Propuesta técnica" not in text


def test_clean_template_contains_no_historical_body() -> None:
    template = RESOURCES / "templates" / "technical-proposal" / "next-standard-v1" / "template.docx"
    with zipfile.ZipFile(template) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    for term in ("OCS", "DDnA", "Self-Serve Data", "Aena T4"):
        assert term not in xml
    for marker in ("[[COVER_LABEL]]", "[[COVER_TITLE]]", "[[INDEX_TITLE]]", "[[INDEX_ENTRIES]]"):
        assert marker in xml


def test_skeleton_content_has_contextual_subsections_and_editorial_guidance() -> None:
    plan = build_document_plan(dossier())
    content = skeleton_content(dossier(), plan)
    section = content["sections"][0]
    rendered_words = sum(len(item.split()) for block in section["blocks"] for item in block["items"])
    assert rendered_words < 100
    assert section["requirements_covered"] == []
    assert section["pending"]
    subheading = next(block for block in section["blocks"] if block["type"] == "subheading")
    assert subheading["text"] != "Ejemplo de subapartado"
    assert subheading["level"] == 2
    assert subheading["id"].startswith(section["id"])
    assert plan["sections"][0]["guidance"]["purpose"]
    assert content["warnings"] == ["Documento esqueleto: no contiene una propuesta técnica redactada."]
def test_existing_run_can_refresh_skeleton_without_provider(tmp_path: Path) -> None:
    dossier_path = tmp_path / "case" / "detailed-analysis" / "run-1" / "output" / "dossier.json"
    dossier_path.parent.mkdir(parents=True)
    dossier_path.write_text(json.dumps(dossier(), ensure_ascii=False), encoding="utf-8")
    run_dir = tmp_path / "generation-run"
    generate_documents(dossier_path, output_dir=run_dir, artifacts=["technical-document"], provider=None)
    content_path = run_dir / "output" / "content.json"
    old_content = json.loads(content_path.read_text(encoding="utf-8"))
    old_content["sections"][0]["blocks"] = []
    content_path.write_text(json.dumps(old_content, ensure_ascii=False), encoding="utf-8")
    result = render_existing_run(run_dir, refresh_skeleton=True)
    refreshed = json.loads(content_path.read_text(encoding="utf-8"))
    assert result["content_refreshed"]
    assert any(block["type"] == "subheading" for block in refreshed["sections"][0]["blocks"])
    assert "Ejemplo de subapartado" not in "\n".join(
        paragraph.text for paragraph in Document(result["document"]).paragraphs
    )


def test_rerender_refreshes_pending_and_rejects_unplanned_coverage(tmp_path: Path) -> None:
    dossier_path = tmp_path / "case" / "detailed-analysis" / "run-1" / "output" / "dossier.json"
    dossier_path.parent.mkdir(parents=True)
    dossier_path.write_text(json.dumps(dossier(), ensure_ascii=False), encoding="utf-8")
    run_dir = tmp_path / "generation-run"
    generate_documents(dossier_path, output_dir=run_dir, artifacts=["technical-document"], provider=None)
    content_path = run_dir / "output" / "content.json"
    content = json.loads(content_path.read_text(encoding="utf-8"))
    content["sections"][0]["pending"] = ["Revisión humana"]
    content_path.write_text(json.dumps(content, ensure_ascii=False), encoding="utf-8")
    render_existing_run(run_dir)
    pending = json.loads((run_dir / "output" / "pending.json").read_text(encoding="utf-8"))
    assert any(item["section_id"] == content["sections"][0]["id"] for item in pending)
    content["sections"][0]["requirements_covered"] = ["No corresponde"]
    content_path.write_text(json.dumps(content, ensure_ascii=False), encoding="utf-8")
    try:
        render_existing_run(run_dir)
    except ValueError as error:
        assert "no respeta" in str(error)
    else:
        raise AssertionError("Debe rechazar cobertura fuera del plan")
