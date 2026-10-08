import json
from pathlib import Path

from docx import Document

from conftest import dossier
from litman_analysis.generation.technical_summary import TechnicalSummaryGenerator
from litman_analysis.generation.technical_summary.renderer import write_technical_summary


def _text(path: Path) -> str:
    doc = Document(path)
    return "\n".join([p.text for p in doc.paragraphs] + [cell.text for table in doc.tables for row in table.rows for cell in row.cells])


def test_technical_summary_contains_dossier_context_and_not_operational_sections(tmp_path: Path) -> None:
    output = tmp_path / "resumen_tecnico.docx"
    write_technical_summary(dossier(), output)
    text = _text(output)
    for value in ("Resumen técnico", "EUR", "Contexto y situación de partida", "Mejorar la calidad", "Trabajos requeridos", "Plan de trabajo", "Entregables del contrato", "Jefe de proyecto", "Criterios que condicionan", "Calidad y limitaciones"):
        assert value in text
    assert "Pendientes para preparar la oferta" not in text
    assert "Documentos que presentar con la oferta" not in text
    assert "None" not in text


def test_technical_summary_rerenders_snapshot_with_its_stable_contract(tmp_path: Path) -> None:
    run = tmp_path / "run"; (run / "input").mkdir(parents=True); (run / "output").mkdir()
    (run / "input" / "dossier.json").write_text(json.dumps(dossier()), encoding="utf-8")
    result = TechnicalSummaryGenerator().rerender(run)
    assert Path(result["documents"][0]).name == "resumen_tecnico.docx"
    assert result["id"] == "technical-summary"
    assert result["generator_version"] == "technical-summary-generator-v1"
    assert result["pending"] == [] and result["metadata"]["format"] == "docx"
